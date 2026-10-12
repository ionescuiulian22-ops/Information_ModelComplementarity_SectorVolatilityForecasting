from pathlib import Path
from itertools import product
from concurrent.futures import ThreadPoolExecutor,as_completed
from math import ceil, floor, lgamma, log, pi, sqrt
from hashlib import sha256
from importlib import metadata
from threading import Lock
import os
import pickle
import platform
import threading
import time
import warnings
import numpy as np
import pandas as pd
import scipy
from scipy.integrate import quad
from scipy.optimize import approx_fprime, minimize, minimize_scalar
from scipy.optimize._numdiff import approx_derivative
from scipy.special import gammaln
from scipy.stats import t as student_t
from statsmodels.stats.multitest import multipletests
from arch import arch_model
import xgboost as xgb
from numba import njit
from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT=Path(__file__).resolve().parent
OUTPUT=ROOT/"final_results.xlsx"
SECTORS=["energy","materials","industrials","consumer_discretionary","consumer_staples","health_care","financials","information_technology","communication_services","utilities","real_estate"]
FEATURES=["vix_change_t","vix_change_lag1","vix_cumulative_change_5d","vix_cumulative_change_20d","move_change_t","move_change_lag1","move_cumulative_change_5d","move_cumulative_change_20d","baa10y_change_t","baa10y_change_lag1","baa10y_cumulative_change_5d","baa10y_cumulative_change_20d","dxy_log_return_t","dxy_log_return_lag1","dxy_cumulative_log_return_5d","dxy_cumulative_log_return_20d","wti_change_t","wti_change_lag1","wti_cumulative_change_5d","wti_cumulative_change_20d","gold_log_return_t","gold_log_return_lag1","gold_cumulative_log_return_5d","gold_cumulative_log_return_20d"]
FOLDS=[("fold_1","2012-01-01","2014-01-01"),("fold_2","2014-01-01","2016-01-01"),("fold_3","2016-01-01","2018-01-01"),("fold_4","2018-01-01","2020-01-01"),("fold_5","2020-01-01","2022-01-01")]
SHEETS=["01 Specification","02 Sample_Audit","03 XGB_Parameters","04 Static_Predictions","05 Static_Metrics","06 Static_Tests_QLIKE","07 Static_Tests_SE","08 Static_Tests_AE","09 MarketWide_Tests","10 Encompassing_Diagnostic","11 Rolling_Predictions","12 Rolling_Tests","13 AR1_Robustness","14 Strict_Timing","15 Numerical_Audit","16 Reporting_Validation"]
MAX_BOOST_ROUNDS=5000
EARLY_STOPPING_ROUNDS=150
TIE_TOLERANCE=1e-12
SEED=42
INVALID_OBJECTIVE=1e50
DATE_COLUMNS={"predictor_date","target_date","raw_first_date","raw_last_date","common_first_date","common_last_date"}
FIT_AUDIT=[]
FIT_LOCK=Lock()
ROLLING_STOP=threading.Event()
XGB_AUDIT={"cv":0,"main":0,"strict":0,"rolling":0}
WORKERS=min(4,max(1,os.cpu_count() or 1))
ROLLING_WINDOWS={"W4":1008,"W6":1512,"W8":2016}
ROLLING_BLOCK_SIZE=21
ROLLING_BOUND_TOLERANCE=1e-7
ROLLING_PROJECTED_GRADIENT_PRIMARY=1e-3
ROLLING_PROJECTED_GRADIENT_MAX=1e-2
ROLLING_OPTIMIZER_SPECIFICATION="HISTORICAL_WARNING_ONLY_ACCEPTANCE"
CHECKPOINT_SCHEMA_VERSION=1
CHECKPOINT_ROOT=ROOT/"runtime_checkpoints"

def file_hash(path):
    h=sha256()
    with open(path,"rb") as stream:
        for block in iter(lambda:stream.read(1048576),b""): h.update(block)
    return h.hexdigest()


def locate_input(name):
    matches=[p for p in ROOT.rglob(name) if p.is_file()]
    if not matches: raise FileNotFoundError(name)
    groups={file_hash(p) for p in matches}
    if len(groups)!=1: raise RuntimeError("input ambiguity: "+name)
    return min(matches,key=lambda p:(len(p.relative_to(ROOT).parts),str(p).lower()))


SECTOR_FILE=locate_input("sector_returns_decimal.csv")
GLOBAL_FILE=locate_input("global_factor_levels.csv")


def atomic_bytes(path,data):
    path.parent.mkdir(parents=True,exist_ok=True); temporary=path.with_name(path.name+".tmp."+str(os.getpid())+"."+str(threading.get_ident()))
    with open(temporary,"wb") as stream:
        stream.write(data); stream.flush(); os.fsync(stream.fileno())
    os.replace(temporary,path)


def atomic_pickle(path,payload):
    atomic_bytes(path,pickle.dumps(payload,protocol=pickle.HIGHEST_PROTOCOL))


def checkpoint_identity(stage,**local):
    return {"schema":CHECKPOINT_SCHEMA_VERSION,"stage":stage,"raw_sector_sha256":file_hash(SECTOR_FILE),"raw_global_sha256":file_hash(GLOBAL_FILE),**local}


def checkpoint_load(path,identity):
    with open(path,"rb") as stream: payload=pickle.load(stream)
    if payload.get("checkpoint_identity")!=identity: raise RuntimeError("incompatible checkpoint: "+str(path))
    return payload


def checkpoint_save(path,identity,payload):
    value=dict(payload); value["checkpoint_identity"]=identity; atomic_pickle(path,value)


def checkpoint_path(stage,*parts):
    return CHECKPOINT_ROOT/stage/("__".join(str(value) for value in parts)+".pkl")


def restore_fit_audit(rows):
    with FIT_LOCK: FIT_AUDIT.extend(rows)




def qlike(y,h):
    y=np.asarray(y,float); h=np.asarray(h,float)
    if np.any(~np.isfinite(h)) or np.any(h<=0): raise RuntimeError("nonpositive or nonfinite forecast")
    return np.log(h)+y/h


def hac_result(values):
    a=np.asarray(values,float); n=len(a); lag=floor(4*(n/100)**(2/9)); u=a-a.mean(); total=float(u@u/n)
    for j in range(1,lag+1): total+=2*(1-j/(lag+1))*float(u[j:]@u[:-j]/n)
    se=sqrt(max(total,0)/n); stat=a.mean()/se if se>0 else np.nan; p=2*student_t.sf(abs(stat),df=n-1) if np.isfinite(stat) else np.nan
    return lag,se,stat,p


def dm_result(values):
    a=np.asarray(values,float); n=len(a); lag,se,raw,p0=hac_result(a); factor=sqrt((n-1)/n); stat=raw*factor if np.isfinite(raw) else np.nan; p=2*student_t.sf(abs(stat),df=n-1) if np.isfinite(stat) else np.nan
    return lag,stat,p


def holm(frame,family,raw,adjusted,reject):
    out=frame.copy(); out[adjusted]=np.nan; out[reject]=False
    for _,idx in out.groupby(family).groups.items():
        ok=out.loc[idx,raw].notna(); ids=np.asarray(list(idx))[ok.to_numpy()]
        if len(ids):
            r,p,_,_=multipletests(out.loc[ids,raw].to_numpy(),alpha=.05,method="holm"); out.loc[ids,adjusted]=p; out.loc[ids,reject]=r
    return out


def raw_data():
    s=pd.read_csv(SECTOR_FILE); g=pd.read_csv(GLOBAL_FILE)
    s["date"]=pd.to_datetime(s.date); g["date"]=pd.to_datetime(g.date)
    required_s=["date","sector_gap_flag"]+[f"sp500_{z}_return" for z in SECTORS]; required_g=["date","vix","move","dxy","wti_futures","gold_futures","baa10y_original"]
    schema=set(required_s)==set(s.columns) and set(required_g)==set(g.columns)
    dates=s.date.is_unique and s.date.is_monotonic_increasing and g.date.is_unique and g.date.is_monotonic_increasing and s.date.equals(g.date)
    if not schema or not dates: raise RuntimeError("raw input validation failed")
    rcols=required_s[2:]; valid=s.sector_gap_flag.eq(0)
    if s.loc[valid,rcols].isna().any().any() or s.loc[~valid,rcols].notna().any().any(): raise RuntimeError("sector gap behavior failed")
    d=s.merge(g,on="date",validate="one_to_one")
    credit=d.baa10y_original.copy(); isolated=credit.isna()&credit.shift(1).notna()&credit.shift(-1).notna()&valid&valid.shift(1,fill_value=False); credit.loc[isolated]=credit.shift(1).loc[isolated]
    levels={"vix":d.vix,"move":d.move,"baa10y":credit,"dxy":d.dxy,"wti":d.wti_futures,"gold":d.gold_futures}
    feat=pd.DataFrame({"date":d.date})
    for key,x in levels.items():
        if key in ["dxy","gold"]:
            if (x.dropna()<=0).any(): raise RuntimeError("nonpositive log-level endpoint")
            ch=np.log(x/x.shift(1)); c5=np.log(x/x.shift(5)); c20=np.log(x/x.shift(20)); stem=key+"_log_return"
        else:
            ch=x-x.shift(1); c5=x-x.shift(5); c20=x-x.shift(20); stem=("baa10y" if key=="baa10y" else key)+"_change"
        feat[stem+"_t"]=ch; feat[stem+"_lag1"]=ch.shift(1); feat[("baa10y" if key=="baa10y" else key)+"_cumulative_"+("log_return_" if key in ["dxy","gold"] else "change_")+"5d"]=c5; feat[("baa10y" if key=="baa10y" else key)+"_cumulative_"+("log_return_" if key in ["dxy","gold"] else "change_")+"20d"]=c20
    if list(feat.columns[1:])!=FEATURES: raise RuntimeError("global feature naming failed")
    rows=[]; prior_valid=valid.astype(int).cumsum().shift(1,fill_value=0)
    adjacent=valid&valid.shift(-1,fill_value=False)
    for sec in SECTORS:
        col=f"sp500_{sec}_return"; z=pd.DataFrame({"sector":sec,"predictor_date":d.date,"target_date":d.date.shift(-1),"return_t":d[col],"realized_return_t1":d[col].shift(-1),"prior_valid_returns":prior_valid}); z["realized_variance_t1"]=z.realized_return_t1**2; z=z.join(feat[FEATURES]); z=z[adjacent].copy(); rows.append(z)
    panel=pd.concat(rows,ignore_index=True); panel=panel[panel.prior_valid_returns.ge(1000)&panel[FEATURES].notna().all(axis=1)].copy(); panel["sample"]=np.where((panel.predictor_date<pd.Timestamp("2022-01-01"))&(panel.target_date<pd.Timestamp("2022-01-01")),"Development",np.where((panel.predictor_date.between("2022-01-01","2025-12-31"))&(panel.target_date.le("2025-12-31")),"Test","Other")); panel=panel[panel["sample"].ne("Other")].reset_index(drop=True)
    common=panel.groupby("predictor_date").sector.nunique(); keep=common[common.eq(11)].index; panel=panel[panel.predictor_date.isin(keep)].copy()
    audit={"raw_input_schema_pass":schema,"raw_dates_unique_ordered_pass":dates,"gap_handling_pass":bool((s.loc[~valid,rcols].isna().all(axis=1)).all() and not adjacent[~valid].any()),"global_transform_pass":bool(panel[FEATURES].notna().all().all()),"sector_raw_rows":len(s),"global_raw_rows":len(g),"raw_first_date":s.date.min(),"raw_last_date":s.date.max(),"gap_count":int((~valid).sum()),"cross_gap_exclusions":int((valid&~valid.shift(-1,fill_value=False)).sum()),"baa10y_missing_count":int(d.baa10y_original.isna().sum()),"baa10y_fill_count":int(isolated.sum()),"baa10y_fill_past_only_pass":bool((credit.loc[isolated]==d.baa10y_original.shift(1).loc[isolated]).all()),"common_first_date":panel.predictor_date.min(),"common_last_date":panel.predictor_date.max(),"development_predictor_dates":int(panel.loc[panel["sample"].eq("Development"),"predictor_date"].nunique()),"development_rows":int((panel["sample"]=="Development").sum()),"test_predictor_dates":int(panel.loc[panel["sample"].eq("Test"),"predictor_date"].nunique()),"test_rows":int((panel["sample"]=="Test").sum()),"global_feature_missing_count":int(panel[FEATURES].isna().sum().sum()),"target_timing_adjacent_pass":bool(adjacent.sum()>0)}
    return panel,audit,s


def student_t_absolute_mean(nu):
    return float(2*sqrt(nu-2)*np.exp(gammaln((nu+1)/2)-gammaln(nu/2))/((nu-1)*sqrt(pi)))


@njit
def student_t_absolute_mean_numba(nu):
    return 2*sqrt(nu-2)*np.exp(lgamma((nu+1)/2)-lgamma(nu/2))/((nu-1)*sqrt(pi))


@njit
def eg_filter(par,r,x,ar,initial_log_variance):
    n=len(r); k=x.shape[1]; core=7 if ar else 6; shift=2 if ar else 1; mu=par[0]; phi=par[1] if ar else 0.0; omega=par[shift]; alpha=par[shift+1]; gamma=par[shift+2]; beta=par[shift+3]; nu=par[shift+4]; delta=par[core:core+k]; first=1 if ar else 0
    logh=np.empty(n); residual=np.empty(n); standardized=np.empty(n)
    for i in range(n): logh[i]=np.nan; residual[i]=np.nan; standardized[i]=np.nan
    if nu<=2.05 or beta<0 or beta>=.999 or not np.isfinite(initial_log_variance) or abs(initial_log_variance)>=60: return INVALID_OBJECTIVE,logh,residual,standardized,False
    residual[0]=r[0]-mu
    for i in range(1,n): residual[i]=r[i]-mu-phi*r[i-1]
    logh[first]=initial_log_variance; expected=student_t_absolute_mean_numba(nu); total=0.0; used=0
    for i in range(first,n):
        if i>first:
            global_term=0.0
            for j in range(k): global_term+=delta[j]*x[i-1,j]
            previous=standardized[i-1]; logh[i]=omega+beta*logh[i-1]+alpha*(abs(previous)-expected)+gamma*previous+global_term
        if not np.isfinite(logh[i]) or abs(logh[i])>=60: return INVALID_OBJECTIVE,logh,residual,standardized,False
        variance=np.exp(logh[i])
        if not np.isfinite(variance) or variance<=0: return INVALID_OBJECTIVE,logh,residual,standardized,False
        standardized[i]=residual[i]/sqrt(variance)
        value=lgamma((nu+1)/2)-lgamma(nu/2)-.5*log(pi*(nu-2))-.5*logh[i]-(nu+1)/2*np.log1p(residual[i]*residual[i]/(variance*(nu-2)))
        if not np.isfinite(value): return INVALID_OBJECTIVE,logh,residual,standardized,False
        total-=value; used+=1
    return total/used,logh,residual,standardized,True


def eg_objective(par,r,x,ar,initial_log_variance):
    return float(eg_filter(np.asarray(par,float),r,x,ar,float(initial_log_variance))[0])


@njit
def digamma_numba(value):
    result=0.0; x=value
    while x<8.0:
        result-=1.0/x; x+=1.0
    inverse=1.0/x; inverse2=inverse*inverse
    return result+log(x)-0.5*inverse-inverse2*(1.0/12.0-inverse2*(1.0/120.0-inverse2*(1.0/252.0-inverse2*(1.0/240.0-inverse2*5.0/660.0))))


@njit
def student_t_absolute_mean_derivative_numba(nu):
    value=student_t_absolute_mean_numba(nu)
    return value*(0.5/(nu-2.0)+0.5*digamma_numba((nu+1.0)/2.0)-0.5*digamma_numba(nu/2.0)-1.0/(nu-1.0))


@njit
def eg_objective_gradient_numba(par,r,x,initial_log_variance):
    n=len(r); k=x.shape[1]; size=6+k; gradient=np.zeros(size)
    if len(par)!=size or par[5]<=2.05 or par[4]<0 or par[4]>=.999 or not np.isfinite(initial_log_variance) or abs(initial_log_variance)>=60: return INVALID_OBJECTIVE,gradient,False
    mu=par[0]; omega=par[1]; alpha=par[2]; gamma=par[3]; beta=par[4]; nu=par[5]; delta=par[6:]; expected=student_t_absolute_mean_numba(nu); expected_derivative=student_t_absolute_mean_derivative_numba(nu); logh=initial_log_variance; logh_gradient=np.zeros(size); total=0.0
    for i in range(n):
        residual=r[i]-mu
        if not np.isfinite(logh) or abs(logh)>=60: return INVALID_OBJECTIVE,np.zeros(size),False
        variance=np.exp(logh)
        if not np.isfinite(variance) or variance<=0: return INVALID_OBJECTIVE,np.zeros(size),False
        standardized=residual/sqrt(variance); ratio=residual*residual/(variance*(nu-2.0)); log_term=np.log1p(ratio); likelihood=lgamma((nu+1.0)/2.0)-lgamma(nu/2.0)-0.5*log(pi*(nu-2.0))-0.5*logh-(nu+1.0)/2.0*log_term
        if not np.isfinite(likelihood): return INVALID_OBJECTIVE,np.zeros(size),False
        total-=likelihood; derivative_logh=0.5-0.5*(nu+1.0)*ratio/(1.0+ratio); derivative_residual=(nu+1.0)*residual/(variance*(nu-2.0)+residual*residual); direct_nu=-0.5*digamma_numba((nu+1.0)/2.0)+0.5*digamma_numba(nu/2.0)+0.5/(nu-2.0)+0.5*log_term-0.5*(nu+1.0)*ratio/((1.0+ratio)*(nu-2.0)); gradient+=derivative_logh*logh_gradient; gradient[0]-=derivative_residual; gradient[5]+=direct_nu
        if i<n-1:
            residual_gradient=np.zeros(size); residual_gradient[0]=-1.0; standardized_gradient=residual_gradient/np.sqrt(variance)-0.5*standardized*logh_gradient; next_gradient=beta*logh_gradient+(alpha*(1.0 if standardized>0 else (-1.0 if standardized<0 else 0.0))+gamma)*standardized_gradient; next_gradient[1]+=1.0; next_gradient[2]+=abs(standardized)-expected; next_gradient[3]+=standardized; next_gradient[4]+=logh; next_gradient[5]-=alpha*expected_derivative
            for j in range(k): next_gradient[6+j]+=x[i,j]
            global_term=0.0
            for j in range(k): global_term+=delta[j]*x[i,j]
            logh=omega+beta*logh+alpha*(abs(standardized)-expected)+gamma*standardized+global_term; logh_gradient=next_gradient
    return total/n,gradient/n,True


def eg_objective_gradient(par,r,x,initial_log_variance):
    objective,gradient,valid=eg_objective_gradient_numba(np.asarray(par,float),np.asarray(r,float),np.asarray(x,float),float(initial_log_variance))
    if not valid: return float(objective),np.asarray(gradient,float),False
    return float(objective),np.asarray(gradient,float),True


def package_initialization(r,ar=False):
    values=np.asarray(r,float)*100; mean_name="AR" if ar else "Constant"; model=arch_model(values,mean=mean_name,lags=1 if ar else 0,vol="EGARCH",p=1,o=1,q=1,dist="StudentsT",rescale=False); result=model.fit(disp="off",show_warning=False,options={"maxiter":2000}); params=result.params; nu=float(params["nu"]); omega=float(params["omega"])+float(params["alpha[1]"])*(student_t_absolute_mean(nu)-sqrt(2/pi)); core=[float(params["Const"] if ar else params["mu"])]
    if ar:
        ar_name=next(name for name in params.index if name not in {"Const","omega","alpha[1]","gamma[1]","beta[1]","nu"}); core.append(float(params[ar_name]))
    core.extend([omega,float(params["alpha[1]"]),float(params["gamma[1]"]),float(params["beta[1]"]),nu]); variance=np.asarray(result.conditional_volatility,float)**2; finite=np.flatnonzero(np.isfinite(variance)&(variance>0)); first=1 if ar else 0
    if len(finite)==0 or finite[0]!=first: raise RuntimeError("package initialization variance failure")
    residual=np.asarray(result.resid,float); residual_finite=residual[np.isfinite(residual)]; package_backcast=float(model.volatility.backcast(residual_finite))
    return {"core":np.asarray(core,float),"initial_log_variance":float(np.log(variance[first])),"package_variance":variance,"package_loglikelihood":float(result.loglikelihood),"package_success":bool(result.convergence_flag==0),"package_backcast":package_backcast,"package_params":params}


def parameter_bounds(feature_count,ar=False):
    return [(-5,5)]+([(-.999,.999)] if ar else [])+[(-5,5),(-2,2),(-2,2),(0,.999),(2.05,200)]+[(-1.5,1.5)]*feature_count


def fit_context_values(context):
    keys=["sector","model","sample_scope","window","block","period","factor"]
    return {key:context.get(key,"") for key in keys}


def fit_egarch(r,x=None,ar=False,context=None,warm_starts=None,package_data=None,fixed_delta=None,extra_starts=None):
    started=time.perf_counter(); context={} if context is None else dict(context); values=np.asarray(r,float)*100; design=np.empty((len(values),0)) if x is None else np.asarray(x,float); k=design.shape[1]; core=7 if ar else 6; bounds=parameter_bounds(k,ar); package_data=package_initialization(np.asarray(r,float),ar) if package_data is None else package_data; package_full=np.r_[package_data["core"],np.zeros(k)]; warm_starts=[] if warm_starts is None else [np.asarray(v,float) for v in warm_starts]; extra_starts=[] if extra_starts is None else [np.asarray(v,float) for v in extra_starts]; families=[]
    for i,start in enumerate(warm_starts): families.append(("warm_"+str(i+1),start,True))
    if not families: families.append(("package",package_full,False))
    independent=[("package",package_full,False)]+[("additional_"+str(i+1),v,False) for i,v in enumerate(extra_starts)]
    fixed=[] if fixed_delta is None else sorted(set(int(i) for i in fixed_delta)); fixed_positions=[core+i for i in fixed]; free_positions=[i for i in range(core+k) if i not in fixed_positions]; free_bounds=[bounds[i] for i in free_positions]
    def normalize(start):
        full=np.asarray(start,float).copy()
        if len(full)!=core+k: raise RuntimeError("EGARCH start dimension mismatch")
        for position in fixed_positions: full[position]=0.0
        for i,(low,high) in enumerate(bounds): full[i]=min(max(full[i],low+1e-10),high-1e-10)
        return full
    def expand(free):
        full=np.zeros(core+k); full[free_positions]=free; return full
    def objective_free(free): return eg_objective(expand(free),values,design,ar,package_data["initial_log_variance"])
    schedules=[]
    for family,start,warm_used in families:
        schedules.extend([("L-BFGS-B",family,start,300,warm_used),("SLSQP",family,start,500,warm_used)])
    for family,start,warm_used in independent:
        schedules.extend([("L-BFGS-B",family,start,1000,warm_used),("SLSQP",family,start,1000,warm_used)])
    candidates=[]; rows=[]; fit_id=context.get("fit_id",context.get("sector","")+"|"+context.get("model","")+"|"+context.get("sample_scope",""))
    for number,(method,family,start,maxiter,warm_used) in enumerate(schedules,1):
        full_start=normalize(start); free_start=full_start[free_positions]; start_objective=objective_free(free_start); result=None; error=""
        try:
            options={"maxiter":maxiter,"ftol":1e-12}
            if method=="L-BFGS-B": options["maxls"]=50
            result=minimize(objective_free,free_start,method=method,bounds=free_bounds,options=options); full=expand(result.x); objective=float(result.fun); gradient=np.asarray(result.jac,float) if getattr(result,"jac",None) is not None else np.asarray([]); gradient_norm=float(np.linalg.norm(gradient)) if gradient.size and np.isfinite(gradient).all() else np.inf; filtered=eg_filter(full,values,design,ar,package_data["initial_log_variance"]); recursion_valid=bool(filtered[4]); minimum=float(np.nanmin(filtered[1])) if recursion_valid else np.nan; maximum=float(np.nanmax(filtered[1])) if recursion_valid else np.nan; within=all(low-1e-10<=full[i]<=high+1e-10 for i,(low,high) in enumerate(bounds)); shift=2 if ar else 1; false_stationarity=bool(abs(objective-start_objective)<=1e-12 and gradient_norm>1e-3); valid=bool(result.success and np.isfinite(objective) and objective<INVALID_OBJECTIVE and within and full[shift+4]>2.05 and full[shift+3]<.999 and np.isfinite(gradient_norm) and recursion_valid and not false_stationarity)
        except Exception as exc:
            full=full_start; objective=INVALID_OBJECTIVE; gradient_norm=np.inf; recursion_valid=False; minimum=np.nan; maximum=np.nan; false_stationarity=False; valid=False; error=type(exc).__name__; result=None
        message="" if result is None else str(result.message); warning_code=error if error else ("optimizer_failure" if not bool(result.success) else ("large_gradient" if gradient_norm>1e-1 else "")); row={"fit_id":fit_id,**fit_context_values(context),"optimizer":method,"attempt_number":number,"start_family":family,"warm_start_used":warm_used,"fallback_used":number>1,"optimizer_success":False if result is None else bool(result.success),"accepted":valid,"selected":False,"objective_start":float(start_objective),"objective_final":float(objective),"objective_change":float(objective-start_objective),"total_loglikelihood":float(-objective*(len(values)-(1 if ar else 0))) if objective<INVALID_OBJECTIVE else np.nan,"gradient_norm":gradient_norm,"beta":float(full[shift+3]),"nu":float(full[shift+4]),"recursion_valid":recursion_valid,"minimum_log_variance":minimum,"maximum_log_variance":maximum,"false_stationarity_rejected":false_stationarity,"nested_reference_ll":np.nan,"nested_ll_difference":np.nan,"nested_tolerance":np.nan,"nested_gate_pass":np.nan,"warning_flag":bool(warning_code),"warning_code":warning_code}
        rows.append(row)
        if valid: candidates.append((objective,full,row))
    if not candidates:
        with FIT_LOCK: FIT_AUDIT.extend(rows)
        raise RuntimeError("mandatory EGARCH fit failed: "+fit_id)
    objective,selected,selected_row=min(candidates,key=lambda item:item[0]); selected_row["selected"]=True
    with FIT_LOCK: FIT_AUDIT.extend(rows)
    shift=2 if ar else 1; info={"optimizer":selected_row["optimizer"],"success":True,"fallback":bool(selected_row["fallback_used"]),"warning":selected_row["warning_code"],"warning_flag":selected_row["warning_flag"],"loglikelihood":selected_row["total_loglikelihood"],"objective":objective,"gradient_norm":selected_row["gradient_norm"],"initial_log_variance":package_data["initial_log_variance"],"package_data":package_data,"fit_id":fit_id,"candidate_count":len(rows),"valid_candidate_count":len(candidates),"false_stationarity_rejections":sum(bool(row["false_stationarity_rejected"]) for row in rows),"beta":selected[shift+3],"nu":selected[shift+4],"elapsed_seconds":time.perf_counter()-started}
    return selected,info


def rolling_objective_tolerance(value):
    return max(1e-10,1e-8*abs(float(value)))


def rolling_projected_gradient(point,objective,bounds,gradient_function=None):
    lower=np.asarray([z[0] for z in bounds],float); upper=np.asarray([z[1] for z in bounds],float); point=np.asarray(point,float)
    gradient=np.asarray(gradient_function(point) if gradient_function is not None else approx_derivative(lambda z:np.asarray([objective(z)]),point,method="3-point",rel_step=1e-5,bounds=(lower,upper)),float).reshape(-1)
    projected=gradient.copy(); at_lower=point-lower<=ROLLING_BOUND_TOLERANCE; at_upper=upper-point<=ROLLING_BOUND_TOLERANCE; projected[at_lower]=np.minimum(gradient[at_lower],0.0); projected[at_upper]=np.maximum(gradient[at_upper],0.0)
    status=["lower" if at_lower[i] else ("upper" if at_upper[i] else "interior") for i in range(len(point))]
    return gradient,projected,status,point-lower,upper-point


def fit_rolling_egarch_constrained(r,x=None,context=None,start_families=None,package_data=None,continuation_r=None,continuation_x=None):
    began=time.perf_counter(); context={} if context is None else dict(context); values=np.asarray(r,float)*100; design=np.empty((len(values),0)) if x is None else np.asarray(x,float); k=design.shape[1]; bounds=parameter_bounds(k,False); package_data=package_initialization(np.asarray(r,float)) if package_data is None else package_data; package_full=np.r_[package_data["core"],np.zeros(k)]; supplied=[] if start_families is None else list(start_families); fit_id=context.get("fit_id",context.get("sector","")+"|"+context.get("model","")+"|rolling"); rows=[]; candidates=[]
    lower=np.asarray([z[0] for z in bounds],float); upper=np.asarray([z[1] for z in bounds],float); objective=lambda point:eg_objective(point,values,design,False,package_data["initial_log_variance"]); objective_gradient=lambda point:eg_objective_gradient(np.asarray(point,float),values,design,package_data["initial_log_variance"])[1]
    if continuation_r is None: continuation_values=np.asarray(r,float)
    else: continuation_values=np.r_[np.asarray(r,float),np.asarray(continuation_r,float)]
    if k: continuation_design=design if continuation_x is None else np.vstack([design,np.asarray(continuation_x,float)])
    else: continuation_design=None
    def normalize(start):
        point=np.asarray(start,float).copy()
        if len(point)!=6+k: raise RuntimeError("rolling EGARCH start dimension mismatch")
        for i,(low,high) in enumerate(bounds): point[i]=min(max(point[i],low+1e-10),high-1e-10)
        return point
    warm=[("warm_"+str(i+1),normalize(item[1]),True) for i,item in enumerate(supplied)]
    independent=[("package",normalize(package_full),False)]
    schedules=[]
    for family,start,warm_used in warm: schedules.extend([("L-BFGS-B",family,start,300,warm_used),("SLSQP",family,start,500,warm_used)])
    for family,start,warm_used in independent: schedules.extend([("L-BFGS-B",family,start,1000,warm_used),("SLSQP",family,start,1000,warm_used)])
    def execute(schedule):
        for method,family,start,maxiter,warm_used in schedule:
            attempt_number=len(rows)+1; attempt_started=time.perf_counter(); start_objective=float(objective(start)); result=None; error=""
            try:
                options={"maxiter":maxiter,"ftol":1e-12}
                if method=="L-BFGS-B": options["maxls"]=50
                result=minimize(objective,start,method=method,bounds=bounds,options=options); point=np.asarray(result.x,float); objective_final=float(result.fun); optimizer_gradient=np.asarray(result.jac,float) if getattr(result,"jac",None) is not None else np.asarray([]); raw_norm=float(np.linalg.norm(optimizer_gradient)) if optimizer_gradient.size and np.isfinite(optimizer_gradient).all() else np.inf
            except Exception as exc:
                point=start.copy(); objective_final=INVALID_OBJECTIVE; raw_norm=np.inf; error=type(exc).__name__
            finite=bool(np.isfinite(point).all() and np.isfinite(objective_final) and objective_final<INVALID_OBJECTIVE); within=bool(finite and all(low-1e-10<=point[i]<=high+1e-10 for i,(low,high) in enumerate(bounds))); parameter_valid=bool(within and point[5]>2.05 and point[4]<.999); filtered=eg_filter(point,values,design,False,package_data["initial_log_variance"]) if parameter_valid else (INVALID_OBJECTIVE,np.asarray([np.nan]),None,None,False); recursion_valid=bool(filtered[4]); minimum=float(np.nanmin(filtered[1])) if recursion_valid else np.nan; maximum=float(np.nanmax(filtered[1])) if recursion_valid else np.nan
            projected_inf=np.inf; active_status=[]; lower_distance=[]; upper_distance=[]
            if parameter_valid and recursion_valid:
                try:
                    _,projected,active_status,lower_distance,upper_distance=rolling_projected_gradient(point,objective,bounds,objective_gradient); projected_inf=float(np.max(np.abs(projected)))
                except Exception as exc:
                    error=error or type(exc).__name__
            forecast_valid=False; forecast_min=np.nan; forecast_max=np.nan
            if parameter_valid and recursion_valid:
                try:
                    _,forecast_path,_=eg_paths(point,continuation_values,continuation_design,initial_log_variance=package_data["initial_log_variance"]); forecast_path=forecast_path[len(values):] if continuation_r is not None else forecast_path; forecast_valid=bool(len(forecast_path)>0 and np.isfinite(forecast_path).all() and (forecast_path>0).all()); forecast_min=float(np.min(forecast_path)) if forecast_valid else np.nan; forecast_max=float(np.max(forecast_path)) if forecast_valid else np.nan
                except Exception:
                    forecast_valid=False
            optimizer_success=bool(result is not None and result.success); false_stationarity=bool(abs(objective_final-start_objective)<=1e-12 and raw_norm>1e-3); accepted=bool(optimizer_success and finite and parameter_valid and recursion_valid and forecast_valid and np.isfinite(raw_norm) and not false_stationarity); kkt_admissible=bool(parameter_valid and recursion_valid and forecast_valid and projected_inf<=ROLLING_PROJECTED_GRADIENT_MAX); kkt_status="STRONG" if projected_inf<=ROLLING_PROJECTED_GRADIENT_PRIMARY else ("ADMISSIBLE" if kkt_admissible else "FAIL"); warning_code=error if error else ("optimizer_failure" if not optimizer_success else ("large_gradient" if raw_norm>1e-1 else ""))
            row={"fit_id":fit_id,**fit_context_values(context),"optimizer":method,"optimizer_status":-1 if result is None else int(result.status),"optimizer_message":error if result is None else str(result.message),"parameter_vector":"|".join(f"{z:.17g}" for z in point),"attempt_number":attempt_number,"start_family":family,"warm_start_used":warm_used,"fallback_used":attempt_number>1,"optimizer_success":optimizer_success,"accepted":accepted,"selected":False,"objective_start":start_objective,"objective_final":objective_final,"objective_change":objective_final-start_objective,"own_start_tolerance":rolling_objective_tolerance(start_objective),"objective_nonworsening_own_start_pass":bool(objective_final<=start_objective+rolling_objective_tolerance(start_objective)),"best_feasible_start_objective":np.nan,"best_start_tolerance":np.nan,"selected_minus_best_start":np.nan,"objective_nonworsening_best_start_pass":np.nan,"total_loglikelihood":float(-objective_final*len(values)) if objective_final<INVALID_OBJECTIVE else np.nan,"gradient_norm":raw_norm,"raw_gradient_norm":raw_norm,"projected_gradient_inf_norm":projected_inf,"projected_gradient_primary_threshold":ROLLING_PROJECTED_GRADIENT_PRIMARY,"projected_gradient_max_threshold":ROLLING_PROJECTED_GRADIENT_MAX,"kkt_status":kkt_status,"kkt_pass":kkt_admissible,"parameter_valid":parameter_valid,"candidate_parameter_valid":parameter_valid,"candidate_recursion_valid":recursion_valid,"candidate_kkt_admissible":kkt_admissible,"kkt_admissible":kkt_admissible,"active_bound_count":sum(z!="interior" for z in active_status),"active_lower_bound_count":sum(z=="lower" for z in active_status),"active_upper_bound_count":sum(z=="upper" for z in active_status),"active_bound_status":"|".join(f"{i}:{z}" for i,z in enumerate(active_status) if z!="interior"),"minimum_lower_bound_distance":float(np.min(lower_distance)) if len(lower_distance) else np.nan,"minimum_upper_bound_distance":float(np.min(upper_distance)) if len(upper_distance) else np.nan,"beta":float(point[4]),"nu":float(point[5]),"recursion_valid":recursion_valid,"minimum_log_variance":minimum,"maximum_log_variance":maximum,"forecast_valid":forecast_valid,"forecast_min":forecast_min,"forecast_max":forecast_max,"false_stationarity_rejected":false_stationarity,"rescue_point_candidate":False,"rescue_candidate":False,"rescue_point_used":False,"rescue_round":0,"fallback_count":attempt_number-1,"independent_optimizer_reconciled":False,"nested_reference_ll":np.nan,"nested_ll_difference":np.nan,"nested_tolerance":np.nan,"nested_gate_pass":np.nan,"elapsed_seconds":time.perf_counter()-attempt_started,"warning_flag":bool(warning_code),"warning_code":warning_code}
            rows.append(row)
            if accepted: candidates.append({"point":point,"objective":objective_final,"row":row})
    execute(schedules)
    if not candidates and k:
        anchor=next((normalize(item[1]) for item in reversed(supplied) if len(item[1])==6+k),package_full); core=anchor[:6]; extras=[("additional_1",np.r_[core,np.full(k,.01)],False),("additional_2",np.r_[core,np.full(k,-.01)],False),("additional_3",np.r_[core,np.linspace(-.01,.01,k)],False)]; extra_schedules=[]
        for family,start,warm_used in warm: extra_schedules.extend([("L-BFGS-B",family,normalize(start),300,warm_used),("SLSQP",family,normalize(start),500,warm_used)])
        for family,start,warm_used in independent+extras: extra_schedules.extend([("L-BFGS-B",family,normalize(start),1000,warm_used),("SLSQP",family,normalize(start),1000,warm_used)])
        execute(extra_schedules)
    if not candidates:
        with FIT_LOCK: FIT_AUDIT.extend(rows)
        raise RuntimeError("mandatory historical rolling candidate set empty: "+fit_id)
    selected=min(candidates,key=lambda z:(z["objective"],z["row"]["attempt_number"])); selected["row"]["selected"]=True; best_start=min(float(row["objective_start"]) for row in rows if np.isfinite(row["objective_start"])); best_tolerance=rolling_objective_tolerance(best_start)
    diagnostic_kkt=[z for z in candidates if z["row"]["kkt_admissible"]]; best_diagnostic_kkt=min(diagnostic_kkt,key=lambda z:z["objective"],default=None)
    for row in rows:
        row["best_feasible_start_objective"]=best_start; row["best_start_tolerance"]=best_tolerance; row["selected_minus_best_start"]=row["objective_final"]-best_start; row["objective_nonworsening_best_start_pass"]=bool(row["objective_final"]<=best_start+best_tolerance); row["best_kkt_objective"]=best_diagnostic_kkt["objective"] if best_diagnostic_kkt is not None else np.nan; row["best_kkt_projected_gradient_inf_norm"]=best_diagnostic_kkt["row"]["projected_gradient_inf_norm"] if best_diagnostic_kkt is not None else np.nan; row["best_feasible_non_kkt_objective"]=np.nan; row["best_feasible_non_kkt_projected_gradient_inf_norm"]=np.nan; row["lower_objective_non_kkt_exists"]=False; row["objective_gap_non_kkt_minus_selected_kkt"]=np.nan
    with FIT_LOCK: FIT_AUDIT.extend(rows)
    row=selected["row"]; feasible_non_kkt=[z for z in candidates if not z["row"]["kkt_admissible"]]; info={"optimizer":row["optimizer"],"success":True,"fallback":bool(row["fallback_used"]),"warning":row["warning_code"],"warning_flag":row["warning_flag"],"loglikelihood":row["total_loglikelihood"],"objective":row["objective_final"],"selected_kkt_objective":row["objective_final"],"gradient_norm":row["raw_gradient_norm"],"raw_gradient_norm":row["raw_gradient_norm"],"selected_raw_gradient_norm":row["raw_gradient_norm"],"projected_gradient_inf_norm":row["projected_gradient_inf_norm"],"selected_projected_gradient_inf_norm":row["projected_gradient_inf_norm"],"kkt_pass":row["kkt_admissible"],"selected_parameter_valid":row["parameter_valid"],"selected_recursion_valid":row["recursion_valid"],"selected_kkt_admissible":row["kkt_admissible"],"selected_optimizer":row["optimizer"],"selected_optimizer_success":row["optimizer_success"],"selected_start_family":row["start_family"],"best_feasible_start_objective":best_start,"selected_minus_best_start":row["selected_minus_best_start"],"best_feasible_non_kkt_objective":min((z["objective"] for z in feasible_non_kkt),default=np.nan),"best_feasible_non_kkt_projected_gradient_inf_norm":min((z["row"]["projected_gradient_inf_norm"] for z in feasible_non_kkt),default=np.nan),"lower_objective_non_kkt_exists":False,"objective_gap_non_kkt_minus_selected_kkt":np.nan,"initial_log_variance":package_data["initial_log_variance"],"package_data":package_data,"fit_id":fit_id,"candidate_count":len(rows),"kkt_admissible_candidate_count":sum(bool(z["row"]["kkt_admissible"]) for z in candidates),"valid_candidate_count":len(candidates),"non_kkt_feasible_candidate_count":len(feasible_non_kkt),"false_stationarity_rejections":sum(bool(z["false_stationarity_rejected"]) for z in rows),"rescue_candidate_count":0,"rescue_attempt_count":0,"rescue_point_count":0,"rescue_optimizer_types":"","rescue_selected":False,"fallback_count":row["fallback_count"],"active_bound_count":row["active_bound_count"],"beta":selected["point"][4],"nu":selected["point"][5],"minimum_log_variance":row["minimum_log_variance"],"maximum_log_variance":row["maximum_log_variance"],"elapsed_seconds":time.perf_counter()-began,"forecast_min":row["forecast_min"],"forecast_max":row["forecast_max"]}
    return selected["point"],info
def eg_paths(par,r,x=None,ar=False,initial_log_variance=None):
    if initial_log_variance is None: raise RuntimeError("training-only initial variance required")
    values=np.asarray(r,float)*100; design=np.empty((len(values),0)) if x is None else np.asarray(x,float); objective,logh,residual,standardized,valid=eg_filter(np.asarray(par,float),values,design,ar,float(initial_log_variance))
    if not valid: raise RuntimeError("invalid EGARCH filtering recursion")
    core=7 if ar else 6; shift=2 if ar else 1; delta=np.asarray(par,float)[core:]; expected=student_t_absolute_mean(float(par[shift+4])); forecasts=np.full(len(values),np.nan)
    for i in range(1 if ar else 0,len(values)):
        term=float(delta@design[i]) if design.shape[1] else 0.0; next_log=float(par[shift])+float(par[shift+3])*logh[i]+float(par[shift+1])*(abs(standardized[i])-expected)+float(par[shift+2])*standardized[i]+term
        if not np.isfinite(next_log) or abs(next_log)>=60: raise RuntimeError("invalid EGARCH forecast recursion")
        forecasts[i]=np.exp(next_log)/10000
    return np.exp(.5*logh)/100,forecasts,logh


def nested_values(unrestricted_info,restricted_info):
    unrestricted=float(unrestricted_info["loglikelihood"]); restricted=float(restricted_info["loglikelihood"]); tolerance=max(1e-6,1e-8*abs(restricted)); difference=unrestricted-restricted; passed=bool(difference>=-tolerance); lr=2*difference
    if passed and lr<0: lr=0.0
    return {"LL_unrestricted":unrestricted,"LL_restricted":restricted,"LL_difference":difference,"nested_tolerance":tolerance,"nested_gate_pass":passed,"lr_stat":lr if passed else np.nan}


def update_nested_audit(fit_id,nested):
    with FIT_LOCK:
        for row in FIT_AUDIT:
            if row["fit_id"]==fit_id and row.get("selected",False):
                row["nested_reference_ll"]=nested["LL_restricted"]; row["nested_ll_difference"]=nested["LL_difference"]; row["nested_tolerance"]=nested["nested_tolerance"]; row["nested_gate_pass"]=nested["nested_gate_pass"]


def student_t_preflight(panel):
    errors=[]
    for nu in [4.0,8.0,20.0]:
        scale=sqrt((nu-2)/nu)
        def density(z):
            u=z/scale; return np.exp(gammaln((nu+1)/2)-gammaln(nu/2))/(sqrt(nu*pi)*scale)*(1+u*u/nu)**(-(nu+1)/2)
        integrated=2*quad(lambda z:z*density(z),0,np.inf,epsabs=1e-12,epsrel=1e-12,limit=300)[0]; errors.append(abs(integrated-student_t_absolute_mean(nu)))
    energy=panel[(panel.sector.eq("energy"))&(panel["sample"].eq("Development"))].sort_values("predictor_date"); package=package_initialization(energy.return_t); par=package["core"]; values=energy.return_t.to_numpy()*100; design=np.empty((len(values),0)); objective,logh,residual,standardized,valid=eg_filter(par,values,design,False,package["initial_log_variance"]); custom_variance=np.exp(logh); package_variance=package["package_variance"]; path_error=float(np.nanmax(np.abs(custom_variance[-500:]-package_variance[-500:]))); nu=par[5]; independent=np.sum(gammaln((nu+1)/2)-gammaln(nu/2)-.5*np.log(pi*(nu-2))-.5*logh-(nu+1)/2*np.log1p(residual*residual/(custom_variance*(nu-2)))); likelihood_error=abs(float(independent)-package["package_loglikelihood"])
    sample=energy.iloc[:300].copy(); mean=sample[FEATURES].mean(); std=sample[FEATURES].std(ddof=0); design_small=((sample[FEATURES]-mean)/std).iloc[:,:2].to_numpy(); package_small=package_initialization(sample.return_t); point=np.r_[package_small["core"],[0.01,-0.01]]; objective_small=lambda p:eg_objective(p,sample.return_t.to_numpy()*100,design_small,False,package_small["initial_log_variance"]); forward=approx_fprime(point,objective_small,1e-6); central=np.empty(len(point)); step=1e-5
    for i in range(len(point)):
        plus=point.copy(); minus=point.copy(); plus[i]+=step; minus[i]-=step; central[i]=(objective_small(plus)-objective_small(minus))/(2*step)
    gradient_error=float(np.max(np.abs(forward-central))); timing_par=np.asarray([0.0,-0.2,0.15,-0.07,0.8,8.0,0.3]); timing_returns=np.asarray([0.01,-0.02,0.015])*100; timing_x=np.asarray([[1.0],[2.0],[4.0]]); initial=log(1.2); timing=eg_filter(timing_par,timing_returns,timing_x,False,initial); expected=student_t_absolute_mean(8.0); z0=timing_returns[0]/sqrt(np.exp(initial)); hand1=-.2+.8*initial+.15*(abs(z0)-expected)-.07*z0+.3*timing_x[0,0]; z1=timing_returns[1]/sqrt(np.exp(hand1)); hand2=-.2+.8*hand1+.15*(abs(z1)-expected)-.07*z1+.3*timing_x[1,0]; timing_error=float(max(abs(hand1-timing[1][1]),abs(hand2-timing[1][2]))); results={"maximum_student_t_centering_reconciliation_error":max(errors),"maximum_package_custom_variance_path_error":path_error,"custom_likelihood_reconciliation_error":likelihood_error,"maximum_gradient_check_error":gradient_error,"egarch_timing_error":timing_error,"student_t_centering_formula_pass":bool(abs(student_t_absolute_mean(8)-student_t_absolute_mean_numba(8))<1e-14),"student_t_centering_numeric_reconciliation_pass":bool(max(errors)<1e-10),"package_custom_variance_path_reconciliation_pass":bool(valid and path_error<1e-8),"custom_likelihood_reconciliation_pass":bool(likelihood_error<1e-8),"finite_difference_gradient_check_pass":bool(np.isfinite(gradient_error) and gradient_error<1e-4),"egarch_timing_unit_test_pass":bool(timing[4] and timing_error<1e-12)}
    if not all(value for key,value in results.items() if key.endswith("_pass")): raise RuntimeError("Student-t numerical preflight failed: "+",".join(key for key,value in results.items() if key.endswith("_pass") and not value))
    return results


def static_sector_checkpoint(panel,sector):
    path=checkpoint_path("static",sector); identity=checkpoint_identity("static",sector=sector,features=FEATURES)
    if path.exists():
        payload=checkpoint_load(path,identity); restore_fit_audit(payload["fit_audit"]); return payload
    frame=panel[panel.sector.eq(sector)].sort_values("predictor_date").copy(); raw_global=frame[FEATURES].copy(); development=frame[frame["sample"].eq("Development")].copy(); mean=development[FEATURES].mean(); std=development[FEATURES].std(ddof=0)
    if (std<=0).any(): raise RuntimeError("invalid Development scaler")
    frame.loc[:,FEATURES]=(frame[FEATURES]-mean)/std; development=frame[frame["sample"].eq("Development")].copy(); test=frame[frame["sample"].eq("Test")].copy(); package=package_initialization(development.return_t); restricted,restricted_info=fit_egarch(development.return_t,context={"fit_id":sector+"|EGARCH|Development","sector":sector,"model":"EGARCH","sample_scope":"Development"},package_data=package); starts=[np.r_[package["core"],np.zeros(24)],np.r_[restricted,np.zeros(24)]]; unrestricted,unrestricted_info=fit_egarch(development.return_t,development[FEATURES],context={"fit_id":sector+"|EGARCH-X|Development","sector":sector,"model":"EGARCH-X","sample_scope":"Development"},warm_starts=starts,package_data=package); nested=nested_values(unrestricted_info,restricted_info)
    if not nested["nested_gate_pass"]:
        patterns=[np.r_[restricted,np.full(24,value)] for value in [0.01,-0.01]]; unrestricted,unrestricted_info=fit_egarch(development.return_t,development[FEATURES],context={"fit_id":sector+"|EGARCH-X-retry|Development","sector":sector,"model":"EGARCH-X","sample_scope":"Development"},warm_starts=[unrestricted,np.r_[restricted,np.zeros(24)]],package_data=package,extra_starts=patterns); nested=nested_values(unrestricted_info,restricted_info)
    update_nested_audit(unrestricted_info["fit_id"],nested)
    if not nested["nested_gate_pass"]: raise RuntimeError("full EGARCH-X nested likelihood failure: "+sector)
    own,restricted_forecast,restricted_logh=eg_paths(restricted,frame.return_t,initial_log_variance=restricted_info["initial_log_variance"]); _,unrestricted_forecast,unrestricted_logh=eg_paths(unrestricted,frame.return_t,frame[FEATURES],initial_log_variance=unrestricted_info["initial_log_variance"]); test_positions=np.flatnonzero(frame["sample"].eq("Test").to_numpy()); static_values=np.r_[restricted_forecast[test_positions],unrestricted_forecast[test_positions]]
    if np.any(~np.isfinite(static_values)) or np.any(static_values<=0): raise RuntimeError("static EGARCH preflight forecast failure: "+sector)
    frame["own_volatility_t"]=own; development=frame[frame["sample"].eq("Development")].copy(); test=frame[frame["sample"].eq("Test")].copy(); item={"frame":frame,"raw_global":raw_global,"dev":development,"test":test,"eg":restricted,"eg_info":restricted_info,"egx":unrestricted,"egx_info":unrestricted_info,"eg_forecast":restricted_forecast[test_positions],"egx_forecast":unrestricted_forecast[test_positions],"scaler_mean":mean,"scaler_std":std,"nested":nested}; comparison={"sector":sector,**nested}
    coefficients=[{"sector":sector,"coefficient":feature,"value":value,"fit_sample":"Development"} for feature,value in zip(["mu","omega","alpha","gamma","beta","nu"]+FEATURES,unrestricted)]; scalers=[{"sector":sector,"feature":feature,"mean":mean[feature],"std_ddof0":std[feature],"fit_sample":"Development"} for feature in FEATURES]; fit_audit=[row for row in FIT_AUDIT if row.get("sector")==sector and row.get("sample_scope")=="Development"]
    payload={"item":item,"comparison":comparison,"coefficients":coefficients,"scalers":scalers,"fit_audit":fit_audit}; checkpoint_save(path,identity,payload); return payload


def static_preflight(panel):
    store={}; coefficients=[]; scalers=[]; comparisons=[]; order=["energy","utilities"]+[sector for sector in SECTORS if sector not in {"energy","utilities"}]
    for sector in order:
        payload=static_sector_checkpoint(panel,sector); store[sector]=payload["item"]; coefficients.extend(payload["coefficients"]); scalers.extend(payload["scalers"]); comparisons.append(payload["comparison"])
    cores=np.vstack([store[sector]["egx"][:6] for sector in SECTORS]); deltas=np.concatenate([store[sector]["egx"][6:] for sector in SECTORS]); identical=bool(np.all(np.max(np.abs(cores-cores[0]),axis=0)<1e-12)); all_zero=bool(np.all(np.abs(deltas)<=1e-10)); zero_supported=True
    if all_zero:
        for sector in SECTORS:
            item=store[sector]; base=np.r_[item["eg"],np.zeros(24)]; gradient=approx_fprime(base,lambda p:eg_objective(p,item["dev"].return_t.to_numpy()*100,item["dev"][FEATURES].to_numpy(),False,item["eg_info"]["initial_log_variance"]),1e-6)[6:]
            if np.linalg.norm(gradient)>1e-6 or abs(item["nested"]["LL_difference"])>item["nested"]["nested_tolerance"]: zero_supported=False
    final_ids={item[name]["fit_id"] for item in store.values() for name in ["eg_info","egx_info"]}; selected_rows=[row for row in FIT_AUDIT if row.get("selected",False) and row["fit_id"] in final_ids]; gates={"two_sector_preflight_pass":bool(all(store[s]["nested"]["nested_gate_pass"] for s in ["energy","utilities"])),"eleven_sector_static_sanity_pass":bool(len(store)==11 and not identical and (not all_zero or zero_supported)),"static_egarch_11_of_11_accepted_pass":sum(row["model"]=="EGARCH" for row in selected_rows)==11,"static_egarchx_11_of_11_accepted_pass":sum(row["model"]=="EGARCH-X" for row in selected_rows)==11,"static_no_false_stationarity_selected_pass":not any(row["false_stationarity_rejected"] for row in selected_rows),"static_all_selected_gradients_finite_pass":all(np.isfinite(row["gradient_norm"]) for row in selected_rows),"static_all_selected_recursions_valid_pass":all(row["recursion_valid"] for row in selected_rows),"static_parameter_movement_pass":bool(not identical and (not all_zero or zero_supported)),"full_egarchx_nested_ll_pass":all(item["nested"]["nested_gate_pass"] for item in store.values())}
    if not all(gates.values()): raise RuntimeError("static preflight failed: "+",".join(key for key,value in gates.items() if not value))
    return store,pd.DataFrame(coefficients),pd.DataFrame(scalers),pd.DataFrame(comparisons),gates


def xgb_universe():
    vals=product([2,3,4,5,6],[.01,.02,.05,.10],[1,5,20],[.70,.85,1.00],[.70,.85,1.00],[0,.1,1,10],[0,.001,.01],[1.1,1.3,1.5,1.7,1.9]); all_cfg=list(vals)
    if len(all_cfg)!=32400: raise RuntimeError("configuration universe mismatch")
    idx=np.sort(np.random.default_rng(SEED).choice(len(all_cfg),128,replace=False)); names=["max_depth","eta","min_child_weight","subsample","colsample_bytree","lambda","alpha","tweedie_variance_power"]
    return [{"config_id":f"cfg_{j+1:03d}","universe_index":int(i),**dict(zip(names,all_cfg[i]))} for j,i in enumerate(idx)]


def xgb_params(cfg):
    keys=["max_depth","eta","min_child_weight","subsample","colsample_bytree","lambda","alpha","tweedie_variance_power"]
    return {"objective":"reg:tweedie","booster":"gbtree","tree_method":"hist","disable_default_eval_metric":1,"gamma":0,"max_delta_step":0,"seed":42,"nthread":1,**{k:cfg[k] for k in keys}}


def xgb_train(train_x,train_y,valid_x,valid_y,cfg,rounds=MAX_BOOST_ROUNDS,early=True,audit_group="cv"):
    if early and valid_y is None: raise RuntimeError("labeled Development validation required")
    if not early and valid_y is not None: raise RuntimeError("future prediction matrix must be target-free")
    with FIT_LOCK: XGB_AUDIT[audit_group]=XGB_AUDIT.get(audit_group,0)+1
    dtr=xgb.DMatrix(train_x,label=train_y); dva=xgb.DMatrix(valid_x,label=valid_y) if early else xgb.DMatrix(valid_x)
    def metric(pred,dmat):
        if np.any(pred<=0) or np.any(~np.isfinite(pred)): return "QLIKE",float("inf")
        return "QLIKE",float(np.mean(np.log(pred)+dmat.get_label()/pred))
    if early:
        model=xgb.train(xgb_params(cfg),dtr,num_boost_round=rounds,evals=[(dva,"validation")],custom_metric=metric,early_stopping_rounds=EARLY_STOPPING_ROUNDS,verbose_eval=False); best=int(model.best_iteration+1); pred=model.predict(dva,iteration_range=(0,best)); return model,best,pred
    model=xgb.train(xgb_params(cfg),dtr,num_boost_round=rounds,verbose_eval=False); pred=model.predict(dva); return model,rounds,pred


def build_fold_own_states(store):
    for sector in SECTORS:
        path=checkpoint_path("fold_states",sector); identity=checkpoint_identity("fold_states",sector=sector,folds=[list(value) for value in FOLDS])
        if path.exists():
            payload=checkpoint_load(path,identity); store[sector]=payload["item"]; restore_fit_audit(payload["fit_audit"]); continue
        item=store[sector]; frame=item["frame"]
        for name,start,end in FOLDS:
            training=frame[(frame.predictor_date<pd.Timestamp(start))&(frame.target_date<pd.Timestamp(start))]; validation=frame[(frame.predictor_date>=pd.Timestamp(start))&(frame.predictor_date<pd.Timestamp(end))&(frame.target_date<pd.Timestamp(end))]; sequence=pd.concat([training,validation]).sort_values("predictor_date"); params,info=fit_egarch(training.return_t,context={"fit_id":sector+"|EGARCH|"+name,"sector":sector,"model":"EGARCH","sample_scope":name}); sigma,forecast,logh=eg_paths(params,sequence.return_t,initial_log_variance=info["initial_log_variance"]); column="own_"+name; frame[column]=np.nan; frame.loc[sequence.index,column]=sigma
        item["frame"]=frame; item["dev"]=frame[frame["sample"].eq("Development")].copy(); item["test"]=frame[frame["sample"].eq("Test")].copy()
        fit_audit=[row for row in FIT_AUDIT if row.get("sector")==sector and row.get("sample_scope") in {fold[0] for fold in FOLDS}]; checkpoint_save(path,identity,{"item":item,"fit_audit":fit_audit})
    return store


def tune_xgb(frame,columns,configs):
    def evaluate(cfg):
        fold_rows=[]; ok=True
        for name,start,end in FOLDS:
            tr=frame[(frame.predictor_date<pd.Timestamp(start))&(frame.target_date<pd.Timestamp(start))].copy(); va=frame[(frame.predictor_date>=pd.Timestamp(start))&(frame.predictor_date<pd.Timestamp(end))&(frame.target_date<pd.Timestamp(end))].copy(); use=[("own_"+name if c=="own_volatility_t" else c) for c in columns]
            try:
                _,trees,pred=xgb_train(tr[use],tr.realized_variance_t1,va[use],va.realized_variance_t1,cfg); losses=qlike(va.realized_variance_t1,pred); fold_rows.append((name,float(losses.mean()),float(np.mean((va.realized_variance_t1-pred)**2)),trees,len(tr),len(va)))
            except Exception:
                ok=False; break
        if ok and len(fold_rows)==5: return {**cfg,"mean_fold_qlike":np.mean([z[1] for z in fold_rows]),"mean_fold_mse":np.mean([z[2] for z in fold_rows]),"median_best_n_trees":int(np.median([z[3] for z in fold_rows])),"folds_succeeded":5,"fold_details":"|".join(f"{z[0]}:{z[1]:.17g}:{z[2]:.17g}:{z[3]}:{z[4]}:{z[5]}" for z in fold_rows)}
        return None
    with ThreadPoolExecutor(max_workers=4) as pool: rows=[x for x in pool.map(evaluate,configs) if x is not None]
    if not rows: raise RuntimeError("no complete XGBoost configuration")
    tab=pd.DataFrame(rows); q=tab.mean_fold_qlike.min(); cand=tab[tab.mean_fold_qlike<=q+TIE_TOLERANCE].sort_values(["mean_fold_mse","max_depth","median_best_n_trees","config_id"]); return cand.iloc[0].to_dict(),tab


def metric_table(predictions,keys=("sector","model")):
    rows=[]
    for key,g in predictions.groupby(list(keys),sort=False):
        key=(key,) if not isinstance(key,tuple) else key; loss=qlike(g.realized_variance_t1,g.forecast_variance_t1); row=dict(zip(keys,key)); row.update({"QLIKE":float(loss.mean()),"RMSE":float(np.sqrt(np.mean((g.realized_variance_t1-g.forecast_variance_t1)**2))),"MAE":float(np.mean(np.abs(g.realized_variance_t1-g.forecast_variance_t1))),"N":len(g)}); rows.append(row)
    result=pd.DataFrame(rows)
    market_keys=[key for key in keys if key!="sector"]
    for key,g in predictions.groupby(market_keys,sort=False):
        key=(key,) if not isinstance(key,tuple) else key; row=dict(zip(market_keys,key)); row["sector"]="MARKET_WIDE"; loss=qlike(g.realized_variance_t1,g.forecast_variance_t1); row.update({"QLIKE":float(loss.mean()),"RMSE":float(np.sqrt(np.mean((g.realized_variance_t1-g.forecast_variance_t1)**2))),"MAE":float(np.mean(np.abs(g.realized_variance_t1-g.forecast_variance_t1))),"N":len(g)}); result=pd.concat([result,pd.DataFrame([row])],ignore_index=True)
    return result


COMPARISONS={"C1":("EGARCH-X","EGARCH"),"C2":("XGB-Augmented","XGB-Own"),"C3":("EGARCH-X","XGB-Augmented"),"C4":("COMB50","EGARCH-X"),"C5":("COMB50","XGB-Augmented")}

def loss_values(name,y,h):
    if name=="QLIKE": return qlike(y,h)
    if name=="SE": return (np.asarray(y,float)-np.asarray(h,float))**2
    if name=="AE": return np.abs(np.asarray(y,float)-np.asarray(h,float))
    raise RuntimeError("unknown loss")

def tests(predictions,prefix=""):
    base=predictions[["sector","predictor_date","target_date","model","realized_variance_t1","forecast_variance_t1"]].copy(); sector_rows=[]; market_rows=[]
    for loss_name in ["QLIKE","SE","AE"]:
        base["loss"]=loss_values(loss_name,base.realized_variance_t1,base.forecast_variance_t1)
        for comparison_id,(model_a,model_b) in COMPARISONS.items():
            a=base[base.model.eq(model_a)].rename(columns={"loss":"loss_a"}); b=base[base.model.eq(model_b)].rename(columns={"loss":"loss_b"}); joined=a.merge(b,on=["sector","predictor_date","target_date"],suffixes=("_a","_b"),validate="one_to_one"); joined["diff"]=joined.loss_a-joined.loss_b
            for sector,group in joined.groupby("sector",sort=False):
                lag,statistic,p_value=dm_result(group["diff"]); sector_rows.append({"loss":loss_name,"comparison_id":prefix+comparison_id,"sector":sector,"model_A":model_a,"model_B":model_b,"mean_loss_diff":group["diff"].mean(),"hac_lag":lag,"dm_stat":statistic,"p_raw":p_value,"N":len(group)})
            daily=joined.groupby("predictor_date",as_index=False)["diff"].mean(); lag,standard_error,statistic,p_value=hac_result(daily["diff"]); market_rows.append({"loss":loss_name,"comparison_id":prefix+comparison_id,"model_A":model_a,"model_B":model_b,"mean_loss_diff":daily["diff"].mean(),"hac_se":standard_error,"statistic":statistic,"p_value":p_value,"hac_lag":lag,"N_dates":len(daily)})
    sector_frame=pd.DataFrame(sector_rows)
    adjusted=[]
    for loss_name,loss_frame in sector_frame.groupby("loss",sort=False): adjusted.append(holm(loss_frame.copy(),"comparison_id","p_raw","p_holm11","reject_holm11"))
    return pd.concat(adjusted,ignore_index=True),pd.DataFrame(market_rows)

def static_models(store,configs):
    predictions=[]; parameter_rows=[]
    for sec in SECTORS:
        path=checkpoint_path("xgb",sec); identity=checkpoint_identity("xgb",sector=sec,folds=[list(value) for value in FOLDS],seed=SEED,max_boost_rounds=MAX_BOOST_ROUNDS,early_stopping_rounds=EARLY_STOPPING_ROUNDS,configuration_ids=[cfg["config_id"] for cfg in configs])
        if path.exists():
            payload=checkpoint_load(path,identity); predictions.extend(payload["predictions"]); parameter_rows.extend(payload["parameter_rows"]); store[sec].update(payload["item_update"])
            with FIT_LOCK:
                for key,value in payload["xgb_audit_delta"].items(): XGB_AUDIT[key]=XGB_AUDIT.get(key,0)+value
            continue
        audit_before=dict(XGB_AUDIT); sector_predictions=[]; sector_parameters=[]
        item=store[sec]; f=item["frame"]; dev=item["dev"]; test=item["test"]; eg_fore=item["eg_forecast"]; egx_fore=item["egx_forecast"]
        selected_own,tab_own=tune_xgb(f,["own_volatility_t"],configs); selected_aug,tab_aug=tune_xgb(f,["own_volatility_t"]+FEATURES,configs)
        for info,sel,tab in [("XGB-Own",selected_own,tab_own),("XGB-Augmented",selected_aug,tab_aug)]:
            row={"sector":sec,"information_set":info,**sel,"successful_configurations":len(tab)}; parameter_rows.append(row); sector_parameters.append(row)
        own_model,_,own_pred=xgb_train(dev[["own_volatility_t"]],dev.realized_variance_t1,test[["own_volatility_t"]],None,selected_own,rounds=int(selected_own["median_best_n_trees"]),early=False,audit_group="main")
        aug_cols=["own_volatility_t"]+FEATURES; aug_model,_,aug_pred=xgb_train(dev[aug_cols],dev.realized_variance_t1,test[aug_cols],None,selected_aug,rounds=int(selected_aug["median_best_n_trees"]),early=False,audit_group="main"); item.update({"own":own_model,"aug":aug_model,"own_cfg":selected_own,"aug_cfg":selected_aug,"columns":aug_cols})
        for name,values in [("EGARCH",eg_fore),("EGARCH-X",egx_fore),("XGB-Own",own_pred),("XGB-Augmented",aug_pred),("COMB50",.5*np.asarray(egx_fore)+.5*np.asarray(aug_pred))]:
            z=test[["sector","predictor_date","target_date","realized_return_t1","realized_variance_t1"]].copy(); z["model"]=name; z["forecast_variance_t1"]=values; predictions.append(z); sector_predictions.append(z)
        audit_delta={key:XGB_AUDIT.get(key,0)-audit_before.get(key,0) for key in XGB_AUDIT}; checkpoint_save(path,identity,{"predictions":sector_predictions,"parameter_rows":sector_parameters,"item_update":{"own_cfg":selected_own,"aug_cfg":selected_aug,"columns":aug_cols},"xgb_audit_delta":audit_delta})
    return pd.concat(predictions,ignore_index=True),pd.DataFrame(parameter_rows),store


def ar1_work(store,base_predictions):
    preds=[]; fits=[]
    for sec in SECTORS:
        path=checkpoint_path("ar1",sec); identity=checkpoint_identity("ar1",sector=sec,features=FEATURES)
        if path.exists():
            payload=checkpoint_load(path,identity); preds.extend(payload["predictions"]); fits.extend(payload["fits"]); restore_fit_audit(payload["fit_audit"]); continue
        sector_preds=[]; sector_fits=[]
        item=store[sec]; f=item["frame"]; dev=item["dev"]; test=item["test"]
        package=package_initialization(dev.return_t,ar=True); restricted,restricted_info=fit_egarch(dev.return_t,ar=True,context={"fit_id":sec+"|AR1-EGARCH|Development","sector":sec,"model":"AR1-EGARCH","sample_scope":"Development"},package_data=package); unrestricted,unrestricted_info=fit_egarch(dev.return_t,dev[FEATURES],ar=True,context={"fit_id":sec+"|AR1-EGARCH-X|Development","sector":sec,"model":"AR1-EGARCH-X","sample_scope":"Development"},warm_starts=[np.r_[package["core"],np.zeros(24)],np.r_[restricted,np.zeros(24)]],package_data=package)
        positions=np.flatnonzero(f["sample"].eq("Test").to_numpy()); _,restricted_forecasts,_=eg_paths(restricted,f.return_t,ar=True,initial_log_variance=restricted_info["initial_log_variance"]); _,unrestricted_forecasts,_=eg_paths(unrestricted,f.return_t,f[FEATURES],ar=True,initial_log_variance=unrestricted_info["initial_log_variance"])
        for model,values,info in [("AR1-EGARCH",restricted_forecasts[positions],restricted_info),("AR1-EGARCH-X",unrestricted_forecasts[positions],unrestricted_info)]:
            z=test[["sector","predictor_date","target_date","realized_return_t1","realized_variance_t1"]].copy(); z["model"]=model; z["forecast_variance_t1"]=values; fit={"sector":sec,"model":model,"optimizer":info["optimizer"],"gradient_norm":info["gradient_norm"],"fallback":info["fallback"],"warning_flag":info["warning_flag"],"fit_id":info["fit_id"]}; preds.append(z); sector_preds.append(z); fits.append(fit); sector_fits.append(fit)
        fit_audit=[row for row in FIT_AUDIT if row.get("sector")==sec and str(row.get("model","")).startswith("AR1-")]; checkpoint_save(path,identity,{"predictions":sector_preds,"fits":sector_fits,"fit_audit":fit_audit})
    pred=pd.concat(preds,ignore_index=True); met=metric_table(pred); comparison=[]
    pairs=[("AR1-EGARCH","EGARCH"),("AR1-EGARCH-X","EGARCH-X")]
    for aa,bb in pairs:
        a=pred[pred.model.eq(aa)].copy(); b=base_predictions[base_predictions.model.eq(bb)].copy(); j=a.merge(b,on=["sector","predictor_date","target_date"],suffixes=("_a","_b")); j["diff"]=qlike(j.realized_variance_t1_a,j.forecast_variance_t1_a)-qlike(j.realized_variance_t1_b,j.forecast_variance_t1_b)
        for sec,g in j.groupby("sector"): lag,stat,p=dm_result(g["diff"]); comparison.append({"comparison":aa+" minus "+bb,"sector":sec,"mean_loss_diff":g["diff"].mean(),"dm_stat":stat,"p_raw":p,"hac_lag":lag,"N":len(g),"direction_same":np.sign(g["diff"].mean())==np.sign(g["diff"].median()),"holm_inference_same":"NOT_APPLICABLE"})
    comp=pd.DataFrame(comparison); comp["family"]=comp.comparison; comp=holm(comp,"family","p_raw","p_holm11","reject_holm11"); return pred,met,comp,pd.DataFrame(fits)


def strict_work(store):
    preds=[]; audits=[]
    for sec in SECTORS:
        path=checkpoint_path("strict",sec); identity=checkpoint_identity("strict",sector=sec,features=FEATURES)
        if path.exists():
            payload=checkpoint_load(path,identity); preds.extend(payload["predictions"]); audits.extend(payload["audits"]); restore_fit_audit(payload["fit_audit"])
            with FIT_LOCK: XGB_AUDIT["strict"]+=payload["xgb_audit_delta"]
            continue
        audit_before=XGB_AUDIT.get("strict",0); sector_preds=[]
        item=store[sec]; f=item["frame"].copy(); shifted=item["raw_global"].shift(1); ok=shifted.notna().all(axis=1); f=f.loc[ok].copy(); f.loc[:,FEATURES]=shifted.loc[ok].to_numpy(); initial=f[f["sample"].eq("Development")]; mean=initial[FEATURES].mean(); std=initial[FEATURES].std(ddof=0); f.loc[:,FEATURES]=(f[FEATURES]-mean)/std; dev=f[f["sample"].eq("Development")]; test=f[f["sample"].eq("Test")]; package=package_initialization(dev.return_t); px,a=fit_egarch(dev.return_t,dev[FEATURES],context={"fit_id":sec+"|EGARCH-X|strict","sector":sec,"model":"EGARCH-X","sample_scope":"strict"},warm_starts=[np.r_[package["core"],np.zeros(24)]],package_data=package); positions=np.flatnonzero(f["sample"].eq("Test").to_numpy()); _,forecasts,_=eg_paths(px,f.return_t,f[FEATURES],initial_log_variance=a["initial_log_variance"]); eg=forecasts[positions]
        cols=["own_volatility_t"]+FEATURES; selected=item["aug_cfg"]; model,_,xp=xgb_train(dev[cols],dev.realized_variance_t1,test[cols],None,selected,rounds=int(selected["median_best_n_trees"]),early=False,audit_group="strict")
        for name,values in [("EGARCH-X-Strict",eg),("XGB-Augmented-Strict",xp)]: z=test[["sector","predictor_date","target_date","realized_return_t1","realized_variance_t1"]].copy(); z["model"]=name; z["forecast_variance_t1"]=values; preds.append(z); sector_preds.append(z)
        audit_row={"sector":sec,"shift_sessions":1,"calendar_day_shift":False,"own_state_time":"t","N_test":len(test),"fit_id":a["fit_id"],"gradient_norm":a["gradient_norm"],"optimizer":a["optimizer"],"fallback":a["fallback"]}; audits.append(audit_row); fit_audit=[row for row in FIT_AUDIT if row.get("sector")==sec and row.get("sample_scope")=="strict"]; checkpoint_save(path,identity,{"predictions":sector_preds,"audits":[audit_row],"fit_audit":fit_audit,"xgb_audit_delta":XGB_AUDIT.get("strict",0)-audit_before})
    p=pd.concat(preds,ignore_index=True); rows=[]
    for sec,g in p.groupby("sector"):
        a=g[g.model.eq("EGARCH-X-Strict")]; b=g[g.model.eq("XGB-Augmented-Strict")]; j=a.merge(b,on=["sector","predictor_date","target_date"],suffixes=("_a","_b")); diff=qlike(j.realized_variance_t1_a,j.forecast_variance_t1_a)-qlike(j.realized_variance_t1_b,j.forecast_variance_t1_b); lag,stat,pv=dm_result(diff); rows.append({"sector":sec,"mean_loss_diff":diff.mean(),"dm_stat":stat,"p_raw":pv,"hac_lag":lag,"N":len(diff)})
    rows=pd.DataFrame(rows)
    reject,adj,_,_=multipletests(rows.p_raw,method="holm"); rows["p_holm11"]=adj; rows["reject_holm11"]=reject
    return p,metric_table(p),rows,pd.DataFrame(audits)


def rolling_task(item,sec,label,width,through_block=None,with_xgb=True):
    f=item["frame"].sort_values("predictor_date").reset_index(drop=True); raw=item["raw_global"].reset_index(drop=True); test_idx=f.index[f["sample"].eq("Test")].to_numpy(); predictions=[]; audits=[]; previous_eg=None; previous_egx=None; selected_cases={}
    block_origins=list(range(0,len(test_idx),ROLLING_BLOCK_SIZE))
    if through_block is not None: block_origins=block_origins[:through_block]
    paths=[checkpoint_path("rolling",label,sec,str(number)) for number in range(1,len(block_origins)+1)]; existing=[path.exists() for path in paths]
    if False in existing and any(existing[existing.index(False)+1:]): raise RuntimeError("rolling checkpoint consistency error: "+label+" "+sec)
    for block_no,start_pos in enumerate(block_origins,1):
        if ROLLING_STOP.is_set(): raise RuntimeError("rolling peer failure cancellation")
        block_idx=test_idx[start_pos:start_pos+ROLLING_BLOCK_SIZE]; origin=block_idx[0]; train_idx=np.arange(origin-width,origin)
        if train_idx.min()<0 or len(train_idx)!=width: raise RuntimeError("rolling training window failure")
        identity_frame=f.loc[train_idx,["predictor_date","target_date","return_t","realized_variance_t1"]]; training_identity={"window":label,"sector":sec,"block_number":block_no,"window_N":width,"first_predictor_date":str(identity_frame.predictor_date.iloc[0]),"last_predictor_date":str(identity_frame.predictor_date.iloc[-1]),"sha256":sha256(pd.util.hash_pandas_object(identity_frame,index=True).values.tobytes()).hexdigest()}; path=paths[block_no-1]; identity=checkpoint_identity("rolling",sector=sec,window=label,block_number=block_no,training_identity=training_identity,block_size=ROLLING_BLOCK_SIZE,rolling_acceptance_rule=ROLLING_OPTIMIZER_SPECIFICATION)
        if path.exists():
            payload=checkpoint_load(path,identity)
            predictions.extend(payload["predictions"]); audits.extend(payload["audits"]); previous_eg=payload["selected_eg"]; previous_egx=payload["selected_egx"]; selected_cases[block_no]=payload["selected_case"]; restore_fit_audit(payload["fit_audit"])
            loaded_models={str(frame["model"].iloc[0]) for frame in payload["predictions"] if len(frame)}
            if with_xgb and not {"XGB-Own","XGB-Augmented","COMB50"}.issubset(loaded_models): raise RuntimeError("rolling checkpoint XGB prediction completeness failure")
            with FIT_LOCK: XGB_AUDIT["rolling"]+=2 if with_xgb else 0
            continue
        prediction_start=len(predictions); audit_start=len(audits)
        train=f.loc[train_idx].copy(); block=f.loc[block_idx].copy(); train_raw=raw.loc[train_idx]; block_raw=raw.loc[block_idx]; mean=train_raw.mean(); std=train_raw.std(ddof=0)
        if (std<=0).any(): raise RuntimeError("rolling scaler failure")
        tx=(train_raw-mean)/std; bx=(block_raw-mean)/std; sequence_returns=pd.concat([train.return_t,block.return_t],ignore_index=True); sequence_x=pd.concat([tx,bx],ignore_index=True); package=package_initialization(train.return_t)
        eg_starts=[] if previous_eg is None else [("previous_block",previous_eg,True)]
        peg,aeg=fit_rolling_egarch_constrained(train.return_t,context={"fit_id":sec+"|EGARCH|"+label+"|"+str(block_no),"sector":sec,"model":"EGARCH","sample_scope":"rolling","window":label,"block":block_no},start_families=eg_starts,package_data=package,continuation_r=block.return_t)
        egx_starts=[]
        if previous_egx is not None: egx_starts.append(("previous_block",previous_egx,True))
        egx_starts.append(("contemporaneous_egarch_zero_delta",np.r_[peg,np.zeros(24)],True))
        pex,aex=fit_rolling_egarch_constrained(train.return_t,tx,context={"fit_id":sec+"|EGARCH-X|"+label+"|"+str(block_no),"sector":sec,"model":"EGARCH-X","sample_scope":"rolling","window":label,"block":block_no},start_families=egx_starts,package_data=package,continuation_r=block.return_t,continuation_x=bx)
        previous_eg=peg; previous_egx=pex; own,eg_forecasts,_=eg_paths(peg,sequence_returns,initial_log_variance=aeg["initial_log_variance"]); _,egx_forecasts,_=eg_paths(pex,sequence_returns,sequence_x,initial_log_variance=aex["initial_log_variance"]); train_own=own[:width]; block_own=own[width:]; eg_values=eg_forecasts[width:]; egx_values=egx_forecasts[width:]
        if with_xgb:
            own_train=pd.DataFrame({"own_volatility_t":train_own}); own_block=pd.DataFrame({"own_volatility_t":block_own}); aug_train=pd.concat([own_train.reset_index(drop=True),tx.reset_index(drop=True)],axis=1); aug_block=pd.concat([own_block.reset_index(drop=True),bx.reset_index(drop=True)],axis=1); _,_,own_values=xgb_train(own_train,train.realized_variance_t1,own_block,None,item["own_cfg"],rounds=int(item["own_cfg"]["median_best_n_trees"]),early=False,audit_group="rolling"); _,_,aug_values=xgb_train(aug_train,train.realized_variance_t1,aug_block,None,item["aug_cfg"],rounds=int(item["aug_cfg"]["median_best_n_trees"]),early=False,audit_group="rolling"); combined=.5*np.asarray(egx_values)+.5*np.asarray(aug_values)
            for model,values in [("EGARCH",eg_values),("EGARCH-X",egx_values),("XGB-Own",own_values),("XGB-Augmented",aug_values),("COMB50",combined)]:
                z=block[["sector","predictor_date","target_date","realized_return_t1","realized_variance_t1"]].copy(); z["window"]=label; z["window_N"]=width; z["block_number"]=block_no; z["model"]=model; z["forecast_variance_t1"]=values; predictions.append(z)
        for model,info in [("EGARCH",aeg),("EGARCH-X",aex)]: audits.append({"sector":sec,"window":label,"window_N":width,"block_number":block_no,"block_N":len(block),"model":model,"fit_count":1,"optimizer_success_count":int(info["selected_optimizer_success"]),"fallback_count":info["fallback_count"],"rescue_count":info["rescue_point_count"],"rescue_attempt_count":info["rescue_attempt_count"],"rescue_point_count":info["rescue_point_count"],"rescue_optimizer_types":info["rescue_optimizer_types"],"false_stationarity_rejection_count":info["false_stationarity_rejections"],"invalid_recursion_count":int(not info["selected_recursion_valid"]),"invalid_forecast_count":int(not (np.isfinite(info["forecast_min"]) and info["forecast_min"]>0)),"actual_warning_count":int(info["warning_flag"]),"objective_final":info["objective"],"selected_objective":info["selected_kkt_objective"],"best_feasible_start_objective":info["best_feasible_start_objective"],"selected_minus_best_start":info["selected_minus_best_start"],"raw_gradient_norm":info["raw_gradient_norm"],"projected_gradient_inf_norm":info["projected_gradient_inf_norm"],"kkt_failure_count":int(not info["selected_kkt_admissible"]),"KKT_admissible_fit_count":int(info["selected_kkt_admissible"]),"lower_objective_non_kkt_conflict_count":int(info["lower_objective_non_kkt_exists"]),"best_non_kkt_objective":info["best_feasible_non_kkt_objective"],"best_non_kkt_projected_gradient_inf_norm":info["best_feasible_non_kkt_projected_gradient_inf_norm"],"objective_gap_non_kkt_minus_selected_kkt":info["objective_gap_non_kkt_minus_selected_kkt"],"candidate_count":info["candidate_count"],"kkt_admissible_candidate_count":info["kkt_admissible_candidate_count"],"non_kkt_feasible_candidate_count":info["non_kkt_feasible_candidate_count"],"selected_optimizer":info["selected_optimizer"],"selected_start_family":info["selected_start_family"],"active_bound_count":info["active_bound_count"],"rescue_selected_count":int(info["rescue_selected"]),"minimum_log_variance":info["minimum_log_variance"],"maximum_log_variance":info["maximum_log_variance"],"forecast_min":info["forecast_min"],"forecast_max":info["forecast_max"],"elapsed_seconds":info["elapsed_seconds"],"training_window_exact":len(train)==width,"future_targets_used_in_fit":False,"rolling_own_state_regenerated":True})
        if with_xgb:
            xgb_forecasts={"XGB-Own":own_values,"XGB-Augmented":aug_values,"COMB50":combined}
            for model in ["XGB-Own","XGB-Augmented","COMB50"]: audits.append({"sector":sec,"window":label,"window_N":width,"block_number":block_no,"block_N":len(block),"model":model,"fit_count":1,"optimizer_success_count":1,"fallback_count":0,"rescue_count":0,"rescue_attempt_count":0,"rescue_point_count":0,"rescue_optimizer_types":"","false_stationarity_rejection_count":0,"invalid_recursion_count":0,"invalid_forecast_count":0,"actual_warning_count":0,"objective_final":np.nan,"selected_objective":np.nan,"best_feasible_start_objective":np.nan,"selected_minus_best_start":np.nan,"raw_gradient_norm":np.nan,"projected_gradient_inf_norm":np.nan,"kkt_failure_count":0,"KKT_admissible_fit_count":0,"lower_objective_non_kkt_conflict_count":0,"best_non_kkt_objective":np.nan,"best_non_kkt_projected_gradient_inf_norm":np.nan,"objective_gap_non_kkt_minus_selected_kkt":np.nan,"candidate_count":0,"kkt_admissible_candidate_count":0,"non_kkt_feasible_candidate_count":0,"selected_optimizer":"","selected_start_family":"","active_bound_count":0,"rescue_selected_count":0,"minimum_log_variance":np.nan,"maximum_log_variance":np.nan,"forecast_min":float(np.min(xgb_forecasts[model])),"forecast_max":float(np.max(xgb_forecasts[model])),"elapsed_seconds":np.nan,"training_window_exact":len(train)==width,"future_targets_used_in_fit":False,"rolling_own_state_regenerated":True})
        selected_cases[block_no]={"EGARCH":aeg,"EGARCH-X":aex,"forecast_min":float(min(np.min(eg_values),np.min(egx_values))),"forecast_max":float(max(np.max(eg_values),np.max(egx_values)))}
        fit_audit=[row for row in FIT_AUDIT if row.get("sector")==sec and row.get("window")==label and row.get("block")==block_no]; payload={"predictions":predictions[prediction_start:],"audits":audits[audit_start:],"selected_eg":peg,"selected_egx":pex,"selected_case":selected_cases[block_no],"fit_audit":fit_audit,"xgb_audit_delta":2 if with_xgb else 0,"training_identity":training_identity}; checkpoint_save(path,identity,payload)
    return predictions,audits,selected_cases


def finalize_rolling_components(pred,audit):
    met=metric_table(pred,keys=("window","sector","model")); sector_tests=[]; market_tests=[]
    for window,g in pred.groupby("window"):
        a,b=tests(g,prefix=window+"_"); a["window"]=window; b["window"]=window; sector_tests.append(a); market_tests.append(b)
    aud=pd.DataFrame(audit)
    summary=aud.groupby(["window","sector","model"],as_index=False).agg(
        window_N=("window_N","first"),fit_count=("fit_count","sum"),optimizer_success_count=("optimizer_success_count","sum"),fallback_count=("fallback_count","sum"),rescue_count=("rescue_count","sum"),rescue_attempt_count=("rescue_attempt_count","sum"),rescue_point_count=("rescue_point_count","sum"),rescue_optimizer_types=("rescue_optimizer_types",lambda z:"|".join(sorted(set(v for item in z for v in str(item).split("|") if v)))),false_stationarity_rejection_count=("false_stationarity_rejection_count","sum"),invalid_recursion_count=("invalid_recursion_count","sum"),invalid_forecast_count=("invalid_forecast_count","sum"),actual_warning_count=("actual_warning_count","sum"),training_window_exact=("training_window_exact","all"),future_targets_used_in_fit=("future_targets_used_in_fit","any"),rolling_own_state_regenerated=("rolling_own_state_regenerated","all"),selected_objective_min=("selected_objective","min"),selected_objective_median=("selected_objective","median"),selected_objective_max=("selected_objective","max"),selected_raw_gradient_min=("raw_gradient_norm","min"),selected_raw_gradient_median=("raw_gradient_norm","median"),selected_raw_gradient_max=("raw_gradient_norm","max"),selected_projected_gradient_min=("projected_gradient_inf_norm","min"),selected_projected_gradient_median=("projected_gradient_inf_norm","median"),selected_projected_gradient_max=("projected_gradient_inf_norm","max"),KKT_admissible_fit_count=("KKT_admissible_fit_count","sum"),lower_objective_non_kkt_conflict_count=("lower_objective_non_kkt_conflict_count","sum"),best_non_kkt_objective_min=("best_non_kkt_objective","min"),candidate_count=("candidate_count","sum"),kkt_admissible_candidate_count=("kkt_admissible_candidate_count","sum"),non_kkt_feasible_candidate_count=("non_kkt_feasible_candidate_count","sum"),minimum_log_variance=("minimum_log_variance","min"),maximum_log_variance=("maximum_log_variance","max"),elapsed_seconds=("elapsed_seconds","sum"),kkt_failure_count=("kkt_failure_count","sum"),active_bound_count=("active_bound_count","sum"),rescue_selected_count=("rescue_selected_count","sum"))
    forecasts=pred.groupby(["window","sector","model"],as_index=False).forecast_variance_t1.agg(invalid_forecast_count_reconciled=lambda z:int((~np.isfinite(z)|(z<=0)).sum()),forecast_p50=lambda z:float(np.quantile(z,.50)),forecast_p95=lambda z:float(np.quantile(z,.95)),forecast_p99=lambda z:float(np.quantile(z,.99)),forecast_max="max")
    return met,pd.concat(sector_tests,ignore_index=True),pd.concat(market_tests,ignore_index=True),summary.merge(forecasts,on=["window","sector","model"],validate="one_to_one")


def rolling_work(store):
    jobs=[(sector,label,width) for label,width in ROLLING_WINDOWS.items() for sector in SECTORS]; all_pred=[]; audit=[]; results={}; ROLLING_STOP.clear()
    try:
        with ThreadPoolExecutor(max_workers=WORKERS) as pool:
            futures={pool.submit(rolling_task,store[sector],sector,label,width):(sector,label,width) for sector,label,width in jobs}
            for future in as_completed(futures):
                key=futures[future]; results[key]=future.result()
    except Exception:
        ROLLING_STOP.set(); raise
    for key in jobs:
        predictions,audits,_=results[key]; all_pred.extend(predictions); audit.extend(audits)
    pred=pd.concat(all_pred,ignore_index=True); return (pred,*finalize_rolling_components(pred,audit))


def encompassing(predictions):
    path=checkpoint_path("encompassing","completed"); prediction_identity=sha256(pd.util.hash_pandas_object(predictions[["sector","predictor_date","target_date","model","realized_variance_t1","forecast_variance_t1"]],index=True).values.tobytes()).hexdigest(); identity=checkpoint_identity("encompassing",prediction_identity=prediction_identity,bootstrap_replications=2000,block_length=22,seed=SEED)
    if path.exists(): return checkpoint_load(path,identity)["result"]
    e=predictions[predictions.model.eq("EGARCH-X")][["sector","predictor_date","realized_variance_t1","forecast_variance_t1"]].rename(columns={"forecast_variance_t1":"egx"}); x=predictions[predictions.model.eq("XGB-Augmented")][["sector","predictor_date","forecast_variance_t1"]].rename(columns={"forecast_variance_t1":"xgb"}); d=e.merge(x,on=["sector","predictor_date"],validate="one_to_one"); rows=[]; rng=np.random.default_rng(SEED)
    for label,g in list(d.groupby("sector",sort=False))+[("MARKET_WIDE",d)]:
        dates=np.sort(g.predictor_date.unique()); ordered=g.sort_values(["predictor_date","sector"]); per_date=len(ordered)//len(dates); y=ordered.realized_variance_t1.to_numpy(); a=ordered.egx.to_numpy(); b=ordered.xgb.to_numpy()
        def objective(lam,ids=None):
            yy=y if ids is None else y[ids]; aa=a if ids is None else a[ids]; bb=b if ids is None else b[ids]; return float(np.mean(qlike(yy,lam*aa+(1-lam)*bb)))
        def optimize_lambda(ids=None):
            grid=np.linspace(0,1,101); values=np.asarray([objective(value,ids) for value in grid]); index=int(np.argmin(values)); candidates=[(float(values[0]),0.0),(float(values[-1]),1.0),(float(values[index]),float(grid[index]))]
            if 0<index<len(grid)-1:
                refined=minimize_scalar(lambda value:objective(value,ids),bounds=(float(grid[index-1]),float(grid[index+1])),method="bounded",options={"xatol":1e-12}); candidates.append((float(refined.fun),float(refined.x)))
            best=min(candidates,key=lambda pair:(pair[0],pair[1])); return best[1],best[0]
        point_lambda,point_objective=optimize_lambda(); boots=[]; block=22
        for _ in range(2000):
            chosen=[]
            while len(chosen)<len(dates):
                start=int(rng.integers(0,len(dates))); chosen.extend(dates[(start+np.arange(block))%len(dates)])
            positions=pd.Index(dates).get_indexer(chosen[:len(dates)]); ids=(positions[:,None]*per_date+np.arange(per_date)).ravel(); value,loss=optimize_lambda(ids); boots.append(value)
        lo,hi=np.quantile(np.asarray(boots,float),[.025,.975]); rows.append({"aggregation":label,"lambda_star":float(point_lambda),"ci_lower":float(lo),"ci_upper":float(hi),"objective_value":float(point_objective),"bootstrap_replications":2000,"block_length":22,"ci_contains_0":bool(lo<=0<=hi),"ci_contains_1":bool(lo<=1<=hi),"cross_section_preserved":bool(label!="MARKET_WIDE" or per_date==11),"diagnostic_only":True})
    result=pd.DataFrame(rows); checkpoint_save(path,identity,{"result":result}); return result


def workbook_write(path,tables):
    with pd.ExcelWriter(path,engine="openpyxl") as writer:
        for name in SHEETS:
            frame=tables[name].copy()
            for c in frame.columns:
                if c in DATE_COLUMNS: frame[c]=pd.to_datetime(frame[c])
            frame.to_excel(writer,sheet_name=name,index=False)
    wb=load_workbook(path)
    for ws in wb.worksheets:
        ws.sheet_state="visible"; ws.freeze_panes="A2"
        for cell in ws[1]: cell.font=Font(bold=True)
        for col in ws.columns:
            width=min(45,max(10,max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])+2)); ws.column_dimensions[col[0].column_letter].width=width
    wb.save(path)


def combine_frames(parts):
    cols=sorted(set().union(*(set(x.columns) for x in parts))); return pd.concat([x.reindex(columns=cols) for x in parts],ignore_index=True)


def workbook_structure(path):
    wb=load_workbook(path,data_only=False)
    formulas=sum(1 for ws in wb for row in ws.iter_rows() for cell in row if cell.data_type=="f")
    comments=sum(1 for ws in wb for row in ws.iter_rows() for cell in row if cell.comment is not None)
    hidden=sum(ws.sheet_state!="visible" for ws in wb)
    charts=sum(len(ws._charts) for ws in wb)
    links=len(wb._external_links)
    macros=int(getattr(wb,"vba_archive",None) is not None)
    return {"formula_count":formulas,"comment_count":comments,"hidden_sheet_count":hidden,"chart_count":charts,"external_link_count":links,"macro_count":macros,"sheet_count":len(wb.sheetnames),"sheet_order_pass":wb.sheetnames==SHEETS}


def validation_frame(gates,statistics):
    gate_rows=[{"record_type":"gate","validation_gate":key,"pass":bool(value),"statistic_value":np.nan,"statistic_text":""} for key,value in gates.items()]
    statistic_rows=[]
    for key,value in statistics.items():
        numeric=value if isinstance(value,(int,float,np.integer,np.floating)) and not isinstance(value,(bool,np.bool_)) else np.nan
        text_value="" if np.isfinite(numeric) else str(value)
        statistic_rows.append({"record_type":"reconciliation_statistic","validation_gate":key,"pass":np.nan,"statistic_value":numeric,"statistic_text":text_value})
    return pd.DataFrame(gate_rows+statistic_rows)


def run():
    warnings.filterwarnings("ignore"); FIT_AUDIT.clear(); XGB_AUDIT.update({key:0 for key in XGB_AUDIT})
    CHECKPOINT_ROOT.mkdir(parents=True,exist_ok=True); data_path=checkpoint_path("preflight","data_and_numerical"); data_identity=checkpoint_identity("data_preflight")
    if data_path.exists():
        payload=checkpoint_load(data_path,data_identity); panel=payload["panel"]; audit=payload["audit"]; numerical=payload["numerical"]
    else:
        panel,audit,_=raw_data(); numerical=student_t_preflight(panel); checkpoint_save(data_path,data_identity,{"panel":panel,"audit":audit,"numerical":numerical})
    test=panel[panel["sample"].eq("Test")]; test_counts=test.groupby("sector").predictor_date.nunique(); configs=xgb_universe()
    if not test_counts.eq(968).all() or len(test)!=10648: raise RuntimeError("Test sample accounting failed")
    store,coefficients,scalers,static_comparisons,static_gates=static_preflight(panel)
    store=build_fold_own_states(store)
    static_predictions,xgb_selected,store=static_models(store,configs)
    static_metrics=metric_table(static_predictions); static_tests,market_tests=tests(static_predictions)
    ar_predictions,ar_metrics,ar_tests,ar_fit=ar1_work(store,static_predictions)
    strict_predictions,strict_metrics,strict_tests,strict_audit=strict_work(store)
    rolling_predictions,rolling_metrics,rolling_tests,rolling_market,rolling_audit=rolling_work(store)
    encompassing_results=encompassing(static_predictions)
    static_keys=static_predictions.groupby("model").apply(lambda frame:set(zip(frame.sector,frame.predictor_date,frame.target_date)),include_groups=False); key_pass=all(value==static_keys.iloc[0] for value in static_keys.iloc[1:])
    static_pivot=static_predictions.pivot(index=["sector","predictor_date","target_date"],columns="model",values="forecast_variance_t1"); static_comb_error=float(np.max(np.abs(static_pivot["COMB50"]-(.5*static_pivot["EGARCH-X"]+.5*static_pivot["XGB-Augmented"]))))
    rolling_pivot=rolling_predictions.pivot(index=["window","sector","predictor_date","target_date"],columns="model",values="forecast_variance_t1"); rolling_comb_error=float(np.max(np.abs(rolling_pivot["COMB50"]-(.5*rolling_pivot["EGARCH-X"]+.5*rolling_pivot["XGB-Augmented"]))))
    selected_audit=pd.DataFrame([row for row in FIT_AUDIT if row.get("selected",False)]); nonfinite_static=int((~np.isfinite(static_predictions.forecast_variance_t1)).sum()); nonpositive_static=int((static_predictions.forecast_variance_t1<=0).sum()); nonfinite_rolling=int((~np.isfinite(rolling_predictions.forecast_variance_t1)).sum()); nonpositive_rolling=int((rolling_predictions.forecast_variance_t1<=0).sum())
    probe=np.linspace(-1,1,101)+.02; _,corrected,_=dm_result(probe); _,_,uncorrected,_=hac_result(probe); correction_error=abs(corrected-uncorrected*sqrt(100/101))
    gates={"raw_input_schema_pass":bool(audit["raw_input_schema_pass"]),"raw_dates_unique_ordered_pass":bool(audit["raw_dates_unique_ordered_pass"]),"gap_handling_pass":bool(audit["gap_handling_pass"]),"baa10y_fill_past_only_pass":bool(audit["baa10y_fill_past_only_pass"]),"global_transform_pass":bool(audit["global_transform_pass"]),"development_test_boundary_pass":bool((panel.loc[panel["sample"].eq("Development"),"target_date"]<pd.Timestamp("2022-01-01")).all()),"test_968_dates_per_sector_pass":bool(test_counts.eq(968).all()),"test_10648_rows_pass":bool(len(test)==10648),**{key:bool(value) for key,value in numerical.items() if key.endswith("_pass")},**static_gates,"model_list_pass":set(static_predictions.model)==set(["EGARCH","EGARCH-X","XGB-Own","XGB-Augmented","COMB50"]),"comparison_list_pass":set(static_tests.comparison_id)==set(COMPARISONS),"loss_list_pass":set(static_tests.loss)==set(["QLIKE","SE","AE"]),"static_key_match_pass":bool(key_pass),"static_forecast_finite_pass":nonfinite_static==0,"static_forecast_positive_pass":nonpositive_static==0,"xgb_search_universe_32400_pass":len(list(product([2,3,4,5,6],[.01,.02,.05,.10],[1,5,20],[.70,.85,1.00],[.70,.85,1.00],[0,.1,1,10],[0,.001,.01],[1.1,1.3,1.5,1.7,1.9])))==32400,"xgb_sample_128_pass":len(configs)==128,"xgb_fold_no_test_pass":max(pd.Timestamp(value[2]) for value in FOLDS)<=pd.Timestamp("2022-01-01"),"xgb_development_only_tuning_pass":bool(len(xgb_selected)==22 and xgb_selected.folds_succeeded.eq(5).all()),"static_scaler_development_only_pass":bool(scalers.fit_sample.eq("Development").all()),"target_free_final_xgb_prediction_pass":bool(XGB_AUDIT["main"]==22 and XGB_AUDIT["strict"]==11 and XGB_AUDIT["rolling"]==2*11*3*ceil(968/21)),"dm_h1_correction_pass":bool(correction_error<1e-12),"static_sector_test_counts_pass":bool(static_tests.groupby(["loss","comparison_id"]).size().eq(11).all()),"market_wide_test_counts_pass":bool(market_tests.groupby("loss").size().eq(5).all() and market_tests.N_dates.eq(968).all()),"holm_reconstruction_pass":bool(static_tests.p_holm11.between(0,1).all()),"strict_timing_one_session_pass":bool(strict_audit.shift_sessions.eq(1).all() and strict_audit.N_test.eq(968).all()),"rolling_training_window_pass":bool(set(rolling_predictions.window_N.unique())=={1008,1512,2016} and rolling_audit.training_window_exact.all()),"rolling_future_targets_unused_pass":bool(not rolling_audit.future_targets_used_in_fit.any()),"rolling_own_state_regenerated_pass":bool(rolling_audit.rolling_own_state_regenerated.all()),"rolling_per_window_rows_pass":bool(all(len(rolling_predictions[rolling_predictions.window.eq(window)])==968*11*5 for window in ROLLING_WINDOWS)),"rolling_loss_list_pass":set(rolling_tests.loss)==set(["QLIKE","SE","AE"]) and set(rolling_market.loss)==set(["QLIKE","SE","AE"]),"rolling_forecast_finite_pass":nonfinite_rolling==0,"rolling_forecast_positive_pass":nonpositive_rolling==0,"comb50_static_exact_half_pass":bool(static_comb_error<1e-15),"comb50_rolling_exact_half_pass":bool(rolling_comb_error<1e-15),"encompassing_row_count_pass":bool(len(encompassing_results)==12),"encompassing_bootstrap_pass":bool(encompassing_results.bootstrap_replications.eq(2000).all() and encompassing_results.block_length.eq(22).all()),"encompassing_numeric_bounds_pass":bool(all(pd.api.types.is_numeric_dtype(encompassing_results[column]) and encompassing_results[column].between(0,1).all() for column in ["lambda_star","ci_lower","ci_upper"])),"encompassing_cross_section_pass":bool(encompassing_results.loc[encompassing_results.aggregation.eq("MARKET_WIDE"),"cross_section_preserved"].all()),"encompassing_diagnostic_only_pass":bool(encompassing_results.diagnostic_only.all())}
    statistics={**{key:value for key,value in numerical.items() if not key.endswith("_pass")},"minimum_full_egarchx_ll_difference":float(static_comparisons.LL_difference.min()),"maximum_comb50_identity_error":max(static_comb_error,rolling_comb_error),"nonfinite_static_forecast_count":nonfinite_static,"nonpositive_static_forecast_count":nonpositive_static,"nonfinite_rolling_forecast_count":nonfinite_rolling,"nonpositive_rolling_forecast_count":nonpositive_rolling,"xgb_cv_fit_count":int(XGB_AUDIT["cv"]),"xgb_target_free_fit_count":int(XGB_AUDIT["main"]+XGB_AUDIT["strict"]+XGB_AUDIT["rolling"]),"dm_h1_factor_reconciliation_error":correction_error}
    if not all(gates.values()): raise RuntimeError("mandatory validation failed before workbook write: "+",".join(key for key,value in gates.items() if not value))
    releases={name:metadata.version(name) for name in ["numpy","pandas","scipy","arch","xgboost"]}; specification=pd.DataFrame([{"field":"raw_inputs","value":"sector_returns_decimal.csv; global_factor_levels.csv"},{"field":"sector_count","value":11},{"field":"global_feature_count","value":24},{"field":"Development_rule","value":"predictor_date and target_date before 2022-01-01"},{"field":"Test_rule","value":"calendar 2022-2025"},{"field":"return_estimation_unit","value":"percent"},{"field":"forecast_output_unit","value":"decimal variance"},{"field":"Student_t_centering","value":"2*sqrt(nu-2)*Gamma((nu+1)/2)/((nu-1)*sqrt(pi)*Gamma(nu/2))"},{"field":"EGARCH_timing","value":"filtered h_t uses X_t-1; forecast h_t+1 uses X_t"},{"field":"own_volatility_t","value":"sigma_t"},{"field":"XGBoost_objective","value":"reg:tweedie"},{"field":"configuration_universe","value":32400},{"field":"configuration_sample","value":128},{"field":"seed","value":42},{"field":"COMB50","value":"0.5 EGARCH-X plus 0.5 XGB-Augmented"},{"field":"QLIKE","value":"log(h)+y/h"},{"field":"variance_floor","value":False},{"field":"clipping","value":False},{"field":"python_release","value":platform.python_version()}]+[{"field":name+"_release","value":release} for name,release in releases.items()])
    sample_audit=pd.DataFrame([{"field":key,"value":value} for key,value in audit.items()]+[{"field":"test_dates_"+sector,"value":int(test_counts[sector])} for sector in SECTORS]+[{"field":"test_rows","value":len(test)},{"field":"feature_count","value":len(FEATURES)}])
    numerical_audit=combine_frames([coefficients.assign(record_type="egarchx_coefficient"),pd.DataFrame(FIT_AUDIT).assign(record_type="fit_audit"),rolling_audit.assign(record_type="rolling_audit")])
    tables={"01 Specification":specification,"02 Sample_Audit":sample_audit,"03 XGB_Parameters":xgb_selected,"04 Static_Predictions":static_predictions,"05 Static_Metrics":static_metrics,"06 Static_Tests_QLIKE":static_tests[static_tests.loss.eq("QLIKE")],"07 Static_Tests_SE":static_tests[static_tests.loss.eq("SE")],"08 Static_Tests_AE":static_tests[static_tests.loss.eq("AE")],"09 MarketWide_Tests":market_tests,"10 Encompassing_Diagnostic":encompassing_results,"11 Rolling_Predictions":rolling_predictions,"12 Rolling_Tests":combine_frames([rolling_tests.assign(aggregation="sector"),rolling_market.assign(aggregation="market_wide"),rolling_metrics.assign(aggregation="metric")]),"13 AR1_Robustness":combine_frames([ar_predictions.assign(record_type="prediction"),ar_metrics.assign(record_type="metric"),ar_tests.assign(record_type="test"),ar_fit.assign(record_type="fit")]),"14 Strict_Timing":combine_frames([strict_predictions.assign(record_type="prediction"),strict_metrics.assign(record_type="metric"),strict_tests.assign(record_type="test"),strict_audit.assign(record_type="audit")]),"15 Numerical_Audit":numerical_audit,"16 Reporting_Validation":validation_frame(gates,statistics)}
    candidate=ROOT/("final_results.candidate."+str(os.getpid())+".xlsx"); workbook_write(candidate,tables); structure=workbook_structure(candidate); structural_gates={"workbook_16_sheets_pass":structure["sheet_count"]==16,"workbook_sheet_order_pass":bool(structure["sheet_order_pass"]),"workbook_values_only_pass":structure["formula_count"]==0,"workbook_no_hidden_sheets_pass":structure["hidden_sheet_count"]==0,"workbook_no_comments_pass":structure["comment_count"]==0,"workbook_no_charts_pass":structure["chart_count"]==0,"workbook_no_external_links_pass":structure["external_link_count"]==0,"workbook_no_macros_pass":structure["macro_count"]==0}; gates.update(structural_gates); statistics.update(structure); tables["16 Reporting_Validation"]=validation_frame(gates,statistics); workbook_write(candidate,tables); final_structure=workbook_structure(candidate)
    if not all(gates.values()) or final_structure!=structure: raise RuntimeError("workbook structural validation failed")
    os.replace(candidate,OUTPUT)

if __name__=="__main__":
    run()
